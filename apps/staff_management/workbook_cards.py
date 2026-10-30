"""Scorecards computed exactly as the desk's workbooks compute them.

Two workbooks are built the same way - the back-office one and the sales
one (``Portfolio Management - Q3 2026_scorecards.xlsb`` over
``2026_Actual_Template.xlsx``) - and :data:`SPECS` holds the few places they
differ: which cells the macro types (the sales macro fills C2 name, C3
branch, E2 role and E3 the BBM, then reads K2), the List's columns, and where
Monthly_Summ keeps what the macro recorded. Everything else is the
workbook's own formulas, so what follows for the back office holds for both.

Against the real files: back office September 108 of 108 exact; sales July
72 of 79 and August 76 of 83 before a fix (Excel matches sheet names without
case - the Commercial cards ask for "Trade_Finance_Income" and the template
says "Trade_Finance_income"), 160 of 162 after it. Months before Q3 were run
on the Q1/Q2 workbooks, so they differ - and a stored card shows what the
macro recorded for them.

Two files make a back-office card, and the desk keeps both:

* ``2026_BO_Non_Sales_Scorecards ….xlsb`` - one card sheet per role
  (TELLER_SCORECARD, BOM_SCORECARD, …), the List of who holds which, and the
  Constants the targets cascade from;
* ``2026_BO_Actual_Template.xlsx`` - a sheet per feed (KYC, CASA, NPS, …),
  a row per person and a column per month.

The desk's macro (``explode_scores`` / ``Do_Monthly_Summary``) walks the List
and, for each person, types their NAME into the card sheet's C2 and their
BRANCH into C4, sets the month in Constants, and reads the score in J6. Every
figure on the card is a formula over those two cells:

    PM          typed on the card (or looked up by branch in Constants)
    YTD Target  =F10*('Constants'!$B$2-$M10)          monthly × months
                =F16*(('Constants'!$B$2-$M16)/12)     annual × months/12
                =F11                                   a level
    Actual      =VLOOKUP($C$2, INDIRECT("[template]"&$L10&"!$c:$R"), $O$1, 0)
    % Score     =IF(H10/G10>1.2, 1.2, H10/G10)  - and 20 other shapes
    Weighted    =I10*E10

This module does not re-type those rules. It reads the formulas themselves
out of the .xlsb - shared formulas included, which the usual readers leave as
``#PTG1!`` - and evaluates them with a small Excel evaluator, person by person
and month by month, exactly as the macro does. So next year's workbook, with
its own lines, weights and caps, is next year's cards with no code change, and
the September card for Bob Waigwa is 92.58% because that is what the
workbook's own J6 says when his name is in C2.

Only the cards are evaluated. The List, BMs, Constants and mapping sheets are
read as the values the workbook last saved; the macro does not recompute them
either.

Everything here is pure: :func:`read_scorecards` and :func:`read_actuals`
turn the two files into plain data, :func:`compute` turns that data into
cards. Storing them is :func:`store`.
"""

from __future__ import annotations

import datetime
import math
import os
import re
import struct
import tempfile

#: The macro's own mapping of List role → card sheet (explode_scores,
#: explode_RBO_scores). BOMCSO and CASHTELLER_SCORECARD are still in the
#: workbook but read the 2023/2024 templates and no List role reaches them.
ROLE_SHEETS = {
    "BOM": "BOM_SCORECARD",
    "CSO": "CSO_SCORECARD",
    "CSO_UB": "CSO_UB_SCORECARD",
    "CSM_DB": "CSM_DB_SCORECARD",
    "Cash_Center": "CASH_TELLER_SCORECARD",
    "Teller": "TELLER_SCORECARD",
    "BSA": "BSA_SCORECARD",
    "RBO": "RBO_SCORECARD",
    "TL_RBO": "TL_RBO_SCORECARD",
}

#: The platform role each card sheet is shown under.
ROLE_CODES = {
    "BOM_SCORECARD": "bo_bom",
    "CSO_SCORECARD": "bo_cso",
    "CSO_UB_SCORECARD": "bo_cso_ub",
    "CSM_DB_SCORECARD": "bo_csm_diaspora",
    "CASH_TELLER_SCORECARD": "bo_cash_teller",
    "TELLER_SCORECARD": "bo_teller",
    "BSA_SCORECARD": "bo_bsa",
    "RBO_SCORECARD": "bo_rbo",
    "TL_RBO_SCORECARD": "bo_tl_rbo",
}

TITLES = {
    "BOM_SCORECARD": "Branch Operations Manager",
    "CSO_SCORECARD": "Customer Service",
    "CSO_UB_SCORECARD": "Customer Service - Ultimate Banking",
    "CSM_DB_SCORECARD": "Client Service Manager - Diaspora Banking",
    "CASH_TELLER_SCORECARD": "Cash Centre Teller",
    "TELLER_SCORECARD": "Teller",
    "BSA_SCORECARD": "Branch Service Ambassador",
    "RBO_SCORECARD": "Retail Back Office (Account Opening)",
    "TL_RBO_SCORECARD": "Retail Back Office Team Leader",
}

#: The sales workbook's macro (explode_scores / Do_Monthly_Summary): List
#: role → card sheet, and the platform role each is shown under.
SALES_ROLE_SHEETS = {
    "SME BBC": "Business_Banking", "SME RM": "Business_Banking",
    "SME ARM": "Business_Banking",
    "PB BBC": "Personal_Banking", "PB RM": "Personal_Banking",
    "PB ARM": "Personal_Banking",
    "Ultimate RM": "Ultimate_RM",
    "SNR MGR ULTIMATE": "SNR_MGR_ULTIMATE",
    "DIASPORA ARM": "Diaspora", "DIASPORA RM": "Diaspora",
    "Mortgage Business ARM": "Mortgage_Business",
    "Mortgage Business RM": "Mortgage_Business",
    "COMMERCIAL RM": "Commercial",
    "COMMERCIAL RM- TRADE": "Commercial_trade",
    "HEAD OF COMMERCIAL": "HEAD_COMMERCIAL",
}
SALES_ROLE_CODES = {
    "SME BBC": "sme_bbc", "SME RM": "sme_rm", "SME ARM": "sme_arm",
    "PB BBC": "pb_bbc", "PB RM": "pb_rm", "PB ARM": "pb_arm",
    "Ultimate RM": "ultimate_rm", "SNR MGR ULTIMATE": "snr_mgr_ultimate",
    "DIASPORA ARM": "diaspora_arm", "DIASPORA RM": "diaspora_rm",
    "Mortgage Business ARM": "mortgage_business_arm",
    "Mortgage Business RM": "mortgage_business_rm",
    "COMMERCIAL RM": "commercial_rm",
    "COMMERCIAL RM- TRADE": "commercial_rm_trade",
    "HEAD OF COMMERCIAL": "head_commercial",
}

BACK_OFFICE, SALES, BRANCH_BUSINESS = "back_office", "sales", "branch_business"


class Spec:
    """Where one workbook keeps things. Everything else is its formulas."""

    def __init__(self, **kw):
        self.__dict__.update(kw)


SPECS = {
    BACK_OFFICE: Spec(
        book=BACK_OFFICE,
        label="back-office",
        marker="TELLER_SCORECARD",
        role_sheets=ROLE_SHEETS,
        # List: PF, sales code, name, role, branch
        list_cols={"pf": 0, "code": 1, "name": 2, "role": 3, "branch": 4},
        # The cells the macro types before it reads the score: C2, C4.
        inputs={(1, 2): "name", (3, 2): "branch"},
        score=(5, 9),                                 # J6
        # Constants!B1 is the VLOOKUP column into the template's C:R
        # (January = 5), B2 the months elapsed.
        month_column=(0, 1), month_count=(1, 1), year=(4, 1),
        header="Perspective", ends=("TOTAL SCORE",),
        cols={"perspective": 1, "measure": 2, "initiative": 3, "weight": 4,
              "prior": None, "growth": None, "fy": 5, "ytd": 6, "actual": 7,
              "score": 8, "weighted": 9, "feed": 11},
        # (sheet, header row, sales code column, January's column)
        summaries=(("Monthly_Summ", 0, 1, 6), ("RBO_Monthly_Summ", 0, 1, 6)),
        files=("scorecards", "actuals"),
        source="the back-office scorecards workbook",
    ),
    SALES: Spec(
        book=SALES,
        label="sales",
        marker="Business_Banking",
        role_sheets=SALES_ROLE_SHEETS,
        list_cols={"pf": None, "code": 0, "name": 1, "role": 2, "branch": 3,
                   "bbm": 4},
        # C2 name, C3 branch, E3 the BBM, E2 the role.
        inputs={(1, 2): "name", (2, 2): "branch", (2, 4): "bbm",
                (1, 4): "role"},
        score=(1, 10),                                # K2
        month_column=(4, 1), month_count=(5, 1), year=(0, 1),
        header="KPI", ends=("Reviewee",),
        cols={"perspective": 1, "measure": 2, "initiative": 3, "weight": 4,
              "prior": 5, "growth": 6, "fy": 7, "ytd": 8, "actual": 9,
              "score": 10, "weighted": 11, "feed": 12},
        summaries=(("Monthly_Summ", 4, 0, 4),),
        files=("sales_scorecards", "sales_actuals"),
        source="the sales scorecards workbook",
    ),
    # BB_Scorecards - the branch managers. Its macro (Do_Monthly_Summary)
    # puts everybody on the List through BM_Scorecard: C2 the name, C3 the
    # branch; K2 the score, written to Branch_Monthly_Summary. It reads the
    # SAME actuals template as the sales workbook, at the BM's own rows.
    BRANCH_BUSINESS: Spec(
        book=BRANCH_BUSINESS,
        label="branch business",
        marker="BM_Scorecard",
        role_sheets={}, default_sheet="BM_Scorecard",
        list_cols={"pf": 0, "code": 1, "name": 2, "role": None, "branch": 3,
                   "bbm": 8},
        inputs={(1, 2): "name", (2, 2): "branch", (2, 4): "bbm"},
        score=(1, 10),                                # K2
        month_column=(0, 1), month_count=(1, 1), year=(5, 1),
        header="KPI", ends=("Reviewee",),
        cols={"perspective": 1, "measure": 2, "initiative": 3, "weight": 4,
              "prior": 5, "growth": 6, "fy": 7, "ytd": 8, "actual": 9,
              "score": 10, "weighted": 11, "feed": 13},
        summaries=(("Branch_Monthly_Summary", 2, 1, 5),),
        files=("bb_scorecards", "sales_actuals"),
        source="the branch business scorecards workbook",
        title="Branch Manager", role_code="bbm",
    ),
}

FIRST_MONTH_INDEX = 5


class WorkbookError(Exception):
    """The file is not what it was said to be; the message says why."""


# ─────────────────────────────────────────────────────────────────────────
# Values
# ─────────────────────────────────────────────────────────────────────────

class Err:
    """An Excel error value. Carried, not raised - Excel's IF can step round
    one, and a card with an error on it shows the error."""

    __slots__ = ("code",)

    def __init__(self, code):
        self.code = code

    def __repr__(self):
        return self.code

    def __eq__(self, other):
        return isinstance(other, Err) and other.code == self.code

    def __hash__(self):
        return hash(self.code)


NA, DIV0, VALUE, REF, NAME = (Err("#N/A"), Err("#DIV/0!"), Err("#VALUE!"),
                              Err("#REF!"), Err("#NAME?"))
_ERROR_CODES = {0x00: "#NULL!", 0x07: "#DIV/0!", 0x0F: "#VALUE!",
                0x17: "#REF!", 0x1D: "#NAME?", 0x24: "#NUM!", 0x2A: "#N/A"}


def _plain(value):
    """A stored value back into Python: errors were written as {"e": code}."""
    if isinstance(value, dict) and "e" in value:
        return Err(value["e"])
    return value


def _store_value(value):
    if isinstance(value, Err):
        return {"e": value.code}
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return {"e": "#NUM!"}
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()
    return value


class Rng:
    __slots__ = ("sheet", "r1", "c1", "r2", "c2")

    def __init__(self, sheet, r1, c1, r2, c2):
        self.sheet, self.r1, self.c1, self.r2, self.c2 = sheet, r1, c1, r2, c2


# ─────────────────────────────────────────────────────────────────────────
# Reading the .xlsb: values, and formulas as a token list
# ─────────────────────────────────────────────────────────────────────────

def _signed(value, bits):
    return value - (1 << bits) if value & (1 << (bits - 1)) else value


def _tokens(raw, row, col, workbook, ptgs, Formula):
    """pyxlsb2's parsed tokens → our own JSON-able RPN.

    ``row``/``col`` are the cell the formula sits in: a shared formula's
    references are offsets from it.
    """
    from pyxlsb2.ptgs import function_names

    out = []
    for t in Formula.parse(raw)._tokens:
        name = type(t).__name__
        if isinstance(t, (ptgs.NumberPtg, ptgs.IntegerPtg)):
            out.append(["n", float(t.value)])
        elif isinstance(t, ptgs.StringPtg):
            out.append(["n", t.value])
        elif isinstance(t, ptgs.BooleanPtg):
            out.append(["n", bool(t.value)])
        elif isinstance(t, ptgs.ErrorPtg):
            out.append(["n", {"e": _ERROR_CODES.get(t.value, "#N/A")}])
        elif isinstance(t, ptgs.MissArgPtg):
            out.append(["n", None])
        elif isinstance(t, ptgs.RefNPtg):
            r = row + _signed(t.row, 32) if not t.row_rel else t.row
            c = col + _signed(t.col, 14) if not t.col_rel else t.col
            out.append(["r", None, r, c])
        elif isinstance(t, ptgs.AreaNPtg):
            r1 = row + _signed(t.first_row, 32) if not t.first_row_rel else t.first_row
            r2 = row + _signed(t.last_row, 32) if not t.last_row_rel else t.last_row
            c1 = col + _signed(t.first_col, 14) if not t.first_col_rel else t.first_col
            c2 = col + _signed(t.last_col, 14) if not t.last_col_rel else t.last_col
            out.append(["a", None, r1, c1, r2, c2])
        elif isinstance(t, ptgs.Ref3dPtg):
            out.append(["r", _sheet_of(workbook, t.extern_sheet_idx),
                        t.row, t.col])
        elif isinstance(t, ptgs.Area3dPtg):
            out.append(["a", _sheet_of(workbook, t.extern_sheet_idx),
                        t.first_row, t.first_col, t.last_row, t.last_col])
        elif isinstance(t, ptgs.RefPtg):
            out.append(["r", None, t.row, t.col])
        elif isinstance(t, ptgs.AreaPtg):
            out.append(["a", None, t.first_row, t.first_col, t.last_row,
                        t.last_col])
        elif name in _BINARY:
            out.append(["b", _BINARY[name]])
        elif name in _UNARY:
            out.append(["u", _UNARY[name]])
        elif isinstance(t, ptgs.ParenPtg):
            continue
        elif isinstance(t, ptgs.AttrPtg):
            # tAttrSum is SUM of one argument, written compactly; the rest
            # (space, if/goto jumps, volatile) change nothing evaluated
            # eagerly.
            if t.attr_sum:
                out.append(["f", "SUM", 1])
            continue
        elif isinstance(t, ptgs.FuncPtg):
            entry = function_names.get(t.idx, ("#UNKNOWN",))
            if len(entry) < 2:
                raise WorkbookError(f"function {entry[0]} with no fixed "
                                      f"argument count")
            out.append(["f", entry[0], entry[1]])
        elif isinstance(t, ptgs.FuncVarPtg):
            if t.idx == 255:
                # A newer function (_xlfn.CONCAT): its name is the token
                # pushed first among its arguments.
                out.append(["udf", t.argc])
            else:
                out.append(["f", function_names.get(t.idx, ("#UNKNOWN",))[0],
                            t.argc])
        elif isinstance(t, ptgs.NamePtg):
            # A defined name. The cards use one only in their titles; read
            # as Excel reads a name it cannot resolve.
            out.append(["n", {"e": "#NAME?"}])
        elif isinstance(t, ptgs.NameXPtg):
            try:
                label = t.stringify([], workbook)
            except Exception:  # noqa: BLE001
                label = "#NAME"
            out.append(["name", label])
        else:
            raise WorkbookError(f"formula element {name} is not supported")
    return out


_BINARY = {"AddPtg": "+", "SubstractPtg": "-", "MultiplyPtg": "*",
           "DividePtg": "/", "PowerPtg": "^", "ConcatPtg": "&",
           "LessPtg": "<", "LessEqualPtg": "<=", "EqualPtg": "=",
           "GreaterEqualPtg": ">=", "GreaterPtg": ">", "NotEqualPtg": "<>",
           "RangePtg": ":"}
_UNARY = {"UPlusPtg": "+", "UMinusPtg": "-", "PercentPtg": "%"}


def _sheet_of(workbook, extern_idx):
    _link, first, last = workbook.resolve_extern_sheet_id(extern_idx)
    if first == last and first >= 0:
        return workbook.sheets[first].name
    raise WorkbookError("a reference to another workbook")


def _read_sheet(workbook, name, with_formulas):
    """→ {"grid": [[value, …], …], "f": {"r,c": tokens}}."""
    from pyxlsb2 import ptgs, recordtypes as rt
    from pyxlsb2.formula import Formula
    from pyxlsb2.recordreader import RecordReader
    from pyxlsb2.records import BaseRecord

    class _Raw(BaseRecord):
        @classmethod
        def read(cls, reader, rectype, reclen):
            r = cls()
            r.data = reader.read(reclen)
            return r

    RecordReader._records[427] = _Raw  # BrtShrFmla - the shared formulas
    RecordReader._records[426] = _Raw  # BrtArrFmla - the array formulas

    sheet = workbook.get_sheet_by_name(name)
    sheet._fp.seek(sheet._data_offset, os.SEEK_SET)
    cells, shared, row = {}, [], -1
    arrays = set()
    for rectype, rec in RecordReader(sheet._fp):
        if rectype == rt.ROW_HDR:
            row = rec.r
        elif rectype in (426, 427):
            # The range, then (array formulas only) one flags byte, then the
            # formula. Both are reached from their cells through a PtgExp.
            r1, r2, c1, c2 = struct.unpack_from("<iiii", rec.data, 0)
            at = 17 if rectype == 426 else 16
            cce = struct.unpack_from("<i", rec.data, at)[0]
            shared.append(((r1, r2, c1, c2),
                           rec.data[at + 4:at + 4 + cce]))
            if rectype == 426:
                arrays.add((r1, r2, c1, c2))
        elif rt.CELL_BLANK <= rectype <= rt.FMLA_ERROR and rec.c is not None:
            value = rec.v
            if rectype == rt.CELL_ISST and isinstance(value, int):
                strings = workbook.stringtable._strings
                value = strings[value] if 0 <= value < len(strings) else value
            elif type(value).__name__ == "ErrorValue":
                value = {"e": "#N/A"}
            cells[(row, rec.c)] = (value, rec.f)
        elif rectype == rt.END_SHEET_DATA:
            break

    height = max((r for r, _ in cells), default=-1) + 1
    width = max((c for _, c in cells), default=-1) + 1
    grid = [[None] * width for _ in range(height)]
    formulas, problems = {}, []
    for (r, c), (value, raw) in cells.items():
        grid[r][c] = _store_value(value) if not isinstance(value, dict) else value
        if not with_formulas or not raw:
            continue
        try:
            first = Formula.parse(raw)._tokens
            array = False
            if first and isinstance(first[0], ptgs.ExpPtg):
                found = next(((box, data) for box, data in shared
                              if box[0] <= r <= box[1]
                              and box[2] <= c <= box[3]), None)
                if found is None:
                    raise WorkbookError("its shared formula is missing")
                array = found[0] in arrays
                raw = found[1]
            tokens = _tokens(raw, r, c, workbook, ptgs, Formula)
            # An array formula ({=…}) evaluates its ranges whole.
            formulas[f"{r},{c}"] = [["array"]] + tokens if array else tokens
        except Exception as exc:  # noqa: BLE001
            problems.append(f"{name}!{_a1(r, c)}: {exc}")
    return {"grid": grid, "f": formulas}, problems


def _a1(row, col):
    letters, n = "", col + 1
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{row + 1}"


def read_scorecards(stream):
    """The scorecards .xlsb → plain data (see :func:`compute`)."""
    with tempfile.NamedTemporaryFile(suffix=".xlsb", delete=False) as tmp:
        for chunk in (stream.chunks() if hasattr(stream, "chunks")
                      else [stream.read()]):
            tmp.write(chunk)
        path = tmp.name
    try:
        return _read_scorecards_path(path)
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def _open_xlsb(path):
    """pyxlsb2's workbook, tolerant of defined names it cannot print.

    It renders every defined name's formula on open, and one that points into
    another workbook (the sales workbook has several) raises - taking the
    whole file down with it. A name's text is never needed here.
    """
    import pyxlsb2.workbook as workbook_module
    from pyxlsb2 import open_workbook
    from pyxlsb2.formula import Formula

    class _Tolerant(Formula):
        def stringify(self, workbook):
            try:
                return super().stringify(workbook)
            except Exception:  # noqa: BLE001
                return ""

    workbook_module.Formula = _Tolerant
    return open_workbook(path)


def _read_scorecards_path(path):
    open_workbook = _open_xlsb

    try:
        workbook = open_workbook(path)
    except Exception as exc:  # noqa: BLE001
        raise WorkbookError(f"that is not an .xlsb workbook ({exc})")
    with workbook:
        names = [s.name for s in workbook.sheets]
        spec = next((s for s in SPECS.values()
                     if "List" in names and s.marker in names), None)
        if spec is None:
            raise WorkbookError(
                "this is neither scorecards workbook - it has no List with a "
                "TELLER_SCORECARD (back office) or Business_Banking (sales) "
                "sheet beside it")
        book = {"book": spec.book, "cards": {}, "tables": {}, "problems": []}
        wanted = set()
        for sheet in sorted(set(spec.role_sheets.values())
                            | ({spec.default_sheet}
                               if getattr(spec, "default_sheet", None)
                               else set())):
            if sheet not in names:
                book["problems"].append(f"{sheet} is not in the workbook")
                continue
            data, problems = _read_sheet(workbook, sheet, True)
            book["cards"][sheet] = data
            book["problems"] += problems
            for tokens in data["f"].values():
                for t in tokens:
                    if t[0] in ("r", "a") and t[1]:
                        wanted.add(t[1])
        summaries = {s[0] for s in spec.summaries}
        for sheet in sorted((wanted | {"List", "Constants"} | summaries)
                            - set(book["cards"])):
            if sheet in names:
                book["tables"][sheet] = _read_sheet(workbook, sheet,
                                                    False)[0]["grid"]
    return book


# ─────────────────────────────────────────────────────────────────────────
# Reading the actuals template
# ─────────────────────────────────────────────────────────────────────────

#: The sales template's sheets that the back-office one does not have.
SALES_TEMPLATE_SHEETS = {"Asset_Drawdown", "Deposit_Growth_Retail",
                         "Asset_Growth", "Revenue_Contribution"}


def read_actuals(stream):
    """The actuals template → {sheet: grid}. Every sheet, every month.

    Every column too: the sales template keeps several blocks of twelve
    months side by side on one sheet (C:R, W:AL, AQ:BF, CE:CT), and the
    cards pick the block with INDIRECT.
    """
    from openpyxl import load_workbook

    try:
        book = load_workbook(stream, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise WorkbookError(f"that is not an .xlsx workbook ({exc})")
    out = {}
    for sheet in book.worksheets:
        grid = []
        for row in sheet.iter_rows(values_only=True):
            values = [_store_value(v) for v in row]
            while values and values[-1] in (None, ""):
                values.pop()
            grid.append(values)
        while grid and not grid[-1]:
            grid.pop()
        out[sheet.title] = grid
    if not out.get("KYC") and not out.get("CASA") and not (
            SALES_TEMPLATE_SHEETS & set(out)):
        raise WorkbookError("this is not an actuals template - it has none "
                              "of the KYC, CASA or Asset_Drawdown sheets")
    return out


def actuals_book(actuals):
    """Which workbook's template this is → SALES or BACK_OFFICE."""
    return SALES if SALES_TEMPLATE_SHEETS & set(actuals) else BACK_OFFICE


#: A month is in the template when it carries at least this share of the
#: figures its fullest month does. The October 2026 template has 1,753
#: non-zero figures for September and ONE for October - a stray, not a month.
MONTH_SHARE = 0.25


def template_months(actuals):
    """The months the template has figures for → [1, 2, … 9].

    Counted in non-zero numbers, because several sheets are pre-filled with
    zeros to December.
    """
    counts = [0] * 12
    for grid in actuals.values():
        for row in grid[1:]:
            for index in range(12):
                column = 6 + index
                value = row[column] if column < len(row) else None
                if isinstance(value, (int, float)) and not isinstance(
                        value, bool) and value != 0:
                    counts[index] += 1
    top = max(counts) or 1
    return [i + 1 for i, n in enumerate(counts)
            if n and n >= top * MONTH_SHARE]


# ─────────────────────────────────────────────────────────────────────────
# Evaluating
# ─────────────────────────────────────────────────────────────────────────

class Evaluator:
    """One card sheet, one person, one month."""

    def __init__(self, book, actuals, holder, month):
        self.spec = SPECS[book.get("book", BACK_OFFICE)]
        self.book, self.actuals, self.sheet = book, actuals, holder.sheet
        self.card = book["cards"][holder.sheet]
        self.month = month
        # What the macro types into the card before it reads the score.
        self.inputs = {cell: getattr(holder, field)
                       for cell, field in self.spec.inputs.items()}
        self.memo, self.busy = {}, set()
        self.array = False

    # -- cells --------------------------------------------------------------
    def value(self, sheet, r, c):
        if sheet in (None, self.sheet):
            return self.cell(r, c)
        if sheet == "Constants" and (r, c) == self.spec.month_column:
            return float(self.month + FIRST_MONTH_INDEX - 1)
        if sheet == "Constants" and (r, c) == self.spec.month_count:
            return float(self.month)
        grid = self.book["tables"].get(sheet)
        if grid is None:
            return REF
        return _plain(_at(grid, r, c))

    def cell(self, r, c):
        key = (r, c)
        if key in self.inputs:
            return self.inputs[key]
        if key in self.memo:
            return self.memo[key]
        tokens = self.card["f"].get(f"{r},{c}")
        if tokens is None:
            return _plain(_at(self.card["grid"], r, c))
        if key in self.busy:
            return REF  # a circular reference
        self.busy.add(key)
        outer = self.array
        try:
            self.array = bool(tokens) and tokens[0] == ["array"]
            result = self.run(tokens[1:] if self.array else tokens, r, c)
            if isinstance(result, Rng):
                result = (self.listed(result) if self.array
                          else self.scalar(result, r, c))
            if isinstance(result, list):
                # A single-cell array formula shows its first element.
                result = result[0] if result else VALUE
        finally:
            self.busy.discard(key)
            self.array = outer
        self.memo[key] = result
        return result

    def listed(self, rng):
        """A range as a flat list - how an array formula sees it."""
        return list(self.cells_of(rng))

    def scalar(self, rng, r, c):
        """Implicit intersection - a range used where one value is wanted."""
        if rng.r1 == rng.r2 and rng.c1 == rng.c2:
            return self.value(rng.sheet, rng.r1, rng.c1)
        if rng.c1 == rng.c2 and rng.r1 <= r <= rng.r2:
            return self.value(rng.sheet, r, rng.c1)
        if rng.r1 == rng.r2 and rng.c1 <= c <= rng.c2:
            return self.value(rng.sheet, rng.r1, c)
        return VALUE

    def cells_of(self, rng):
        if rng.sheet and rng.sheet.startswith("@"):
            grid = self.actuals.get(rng.sheet[1:], [])
            last = min(rng.r2, len(grid) - 1)
            for r in range(rng.r1, last + 1):
                for c in range(rng.c1, rng.c2 + 1):
                    yield _plain(_at(grid, r, c))
            return
        if rng.sheet in (None, self.sheet):
            last = min(rng.r2, len(self.card["grid"]) + 5)
        else:
            last = min(rng.r2, len(self.book["tables"].get(rng.sheet, [])))
        for r in range(rng.r1, last + 1):
            for c in range(rng.c1, rng.c2 + 1):
                yield self.value(rng.sheet, r, c)

    # -- the stack machine -------------------------------------------------
    def run(self, tokens, r, c):
        stack = []
        for t in tokens:
            kind = t[0]
            if kind == "n":
                stack.append(_plain(t[1]))
            elif kind == "r":
                stack.append(Rng(t[1], t[2], t[3], t[2], t[3]))
            elif kind == "a":
                stack.append(Rng(t[1], t[2], t[3], t[4], t[5]))
            elif kind == "name":
                stack.append(("name", t[1]))
            elif kind == "u":
                operand = self.get(stack.pop(), r, c)
                if isinstance(operand, list):
                    stack.append([self.unary(t[1], v, r, c) for v in operand])
                else:
                    stack.append(self.unary(t[1], operand, r, c))
            elif kind == "b":
                right, left = stack.pop(), stack.pop()
                stack.append(self.binary(t[1], left, right, r, c))
            elif kind == "f":
                argc = t[2]
                args = stack[len(stack) - argc:] if argc else []
                del stack[len(stack) - argc:]
                stack.append(self.call(t[1], args, r, c))
            elif kind == "udf":
                argc = t[1]
                args = stack[len(stack) - argc:]
                del stack[len(stack) - argc:]
                label = args[0][1] if isinstance(args[0], tuple) else ""
                stack.append(self.call(label.replace("_xlfn.", "").upper(),
                                       args[1:], r, c))
        return stack[-1] if stack else None

    def unary(self, op, v, r, c):
        v = self.number(v, r, c)
        if isinstance(v, Err):
            return v
        if op == "-":
            return -v
        if op == "%":
            return v / 100.0
        return v

    def get(self, v, r, c):
        if not isinstance(v, Rng):
            return v
        if self.array and not (v.r1 == v.r2 and v.c1 == v.c2):
            return self.listed(v)
        return self.scalar(v, r, c)

    def number(self, v, r, c):
        v = self.get(v, r, c)
        if isinstance(v, Err):
            return v
        if v is None or v == "":
            return 0.0
        if isinstance(v, bool):
            return 1.0 if v else 0.0
        if isinstance(v, (int, float)):
            return float(v)
        try:
            return float(str(v).replace(",", "").strip())
        except ValueError:
            return VALUE

    def text(self, v, r, c):
        v = self.get(v, r, c)
        if isinstance(v, Err):
            return v
        if v is None:
            return ""
        if isinstance(v, bool):
            return "TRUE" if v else "FALSE"
        if isinstance(v, float):
            return _general(v)
        return str(v)

    def binary(self, op, left, right, r, c):
        if op == ":":
            if isinstance(left, Rng) and isinstance(right, Rng):
                return Rng(left.sheet, min(left.r1, right.r1),
                           min(left.c1, right.c1), max(left.r2, right.r2),
                           max(left.c2, right.c2))
            return VALUE
        if self.array:
            left, right = self.get(left, r, c), self.get(right, r, c)
            if isinstance(left, list) or isinstance(right, list):
                # Element by element, a single value standing for every one.
                size = max(len(x) for x in (left, right)
                           if isinstance(x, list))
                lefts = left if isinstance(left, list) else [left] * size
                rights = right if isinstance(right, list) else [right] * size
                if len(lefts) != len(rights):
                    return [Err("#N/A")] * size
                return [self.scalar_binary(op, a, b, r, c)
                        for a, b in zip(lefts, rights)]
        return self.scalar_binary(op, left, right, r, c)

    def scalar_binary(self, op, left, right, r, c):
        if op == "&":
            a, b = self.text(left, r, c), self.text(right, r, c)
            if isinstance(a, Err):
                return a
            if isinstance(b, Err):
                return b
            return a + b
        if op in ("<", "<=", "=", ">=", ">", "<>"):
            a, b = self.get(left, r, c), self.get(right, r, c)
            if isinstance(a, Err):
                return a
            if isinstance(b, Err):
                return b
            return _compare(op, a, b)
        a, b = self.number(left, r, c), self.number(right, r, c)
        if isinstance(a, Err):
            return a
        if isinstance(b, Err):
            return b
        if op == "+":
            return a + b
        if op == "-":
            return a - b
        if op == "*":
            return a * b
        if op == "/":
            return DIV0 if b == 0 else a / b
        if op == "^":
            try:
                return float(a ** b)
            except (OverflowError, ZeroDivisionError, ValueError):
                return Err("#NUM!")
        return VALUE

    # -- functions ----------------------------------------------------------
    def call(self, fn, args, r, c):
        fn = fn.upper()
        if fn == "IF":
            cond = self.get(args[0], r, c)
            if isinstance(cond, Err):
                return cond
            truth = _truth(cond)
            if isinstance(truth, Err):
                return truth
            if truth:
                return self.get(args[1], r, c) if len(args) > 1 else True
            return self.get(args[2], r, c) if len(args) > 2 else False
        if fn in ("AND", "OR"):
            values = []
            for a in args:
                for v in (self.cells_of(a) if isinstance(a, Rng) else [a]):
                    if isinstance(v, Err):
                        return v
                    if v is None or isinstance(v, str):
                        continue
                    values.append(bool(v))
            if not values:
                return VALUE
            return all(values) if fn == "AND" else any(values)
        if fn == "NOT":
            t = _truth(self.get(args[0], r, c))
            return t if isinstance(t, Err) else not t
        if fn in ("MIN", "MAX", "SUM", "AVERAGE", "COUNT"):
            numbers = []
            for a in args:
                if isinstance(a, (Rng, list)):
                    for v in (self.cells_of(a) if isinstance(a, Rng) else a):
                        if isinstance(v, Err):
                            return v
                        if isinstance(v, (int, float)) and not isinstance(
                                v, bool):
                            numbers.append(float(v))
                else:
                    v = self.number(a, r, c)
                    if isinstance(v, Err):
                        return v
                    numbers.append(v)
            if fn == "COUNT":
                return float(len(numbers))
            if fn == "SUM":
                total = 0.0
                for n in numbers:  # left to right, as Excel adds
                    total += n
                return total
            if fn == "AVERAGE":
                return (sum(numbers) / len(numbers)) if numbers else DIV0
            if not numbers:
                return 0.0
            return min(numbers) if fn == "MIN" else max(numbers)
        if fn == "ABS":
            v = self.number(args[0], r, c)
            return v if isinstance(v, Err) else abs(v)
        if fn in ("ROUND", "ROUNDUP", "ROUNDDOWN"):
            v, d = self.number(args[0], r, c), self.number(args[1], r, c)
            if isinstance(v, Err):
                return v
            if isinstance(d, Err):
                return d
            return _round(fn, v, int(d))
        if fn in ("IFERROR",):
            v = self.get(args[0], r, c)
            return self.get(args[1], r, c) if isinstance(v, Err) else v
        if fn in ("ISERROR", "ISNA"):
            v = self.get(args[0], r, c)
            if fn == "ISNA":
                return v == NA
            return isinstance(v, Err)
        if fn == "VLOOKUP":
            return self.vlookup(args, r, c)
        if fn == "INDIRECT":
            return self.indirect(self.text(args[0], r, c))
        if fn in ("CONCAT", "CONCATENATE"):
            out = ""
            for a in args:
                for v in (self.cells_of(a) if isinstance(a, Rng) else [a]):
                    s = self.text(v, r, c)
                    if isinstance(s, Err):
                        return s
                    out += s
            return out
        if fn == "TEXT":
            v = self.get(args[0], r, c)
            fmt = self.text(args[1], r, c)
            return _format(v, fmt)
        if fn == "YEAR":
            v = self.number(args[0], r, c)
            if isinstance(v, Err):
                return v
            return float(_from_serial(v).year)
        if fn == "MONTH":
            v = self.number(args[0], r, c)
            return v if isinstance(v, Err) else float(_from_serial(v).month)
        if fn == "DATE":
            y, m, d = (self.number(a, r, c) for a in args[:3])
            for v in (y, m, d):
                if isinstance(v, Err):
                    return v
            return _to_serial(int(y), int(m), int(d))
        if fn == "COLUMN":
            if args and isinstance(args[0], Rng):
                return float(args[0].c1 + 1)
            return float(c + 1)
        if fn == "ROW":
            if args and isinstance(args[0], Rng):
                rng = args[0]
                if self.array and rng.r2 > rng.r1:
                    return [float(i + 1) for i in range(rng.r1, rng.r2 + 1)]
                return float(rng.r1 + 1)
            return float(r + 1)
        if fn == "SUMIF":
            return self.sumifs(args[2] if len(args) > 2 else args[0],
                               [(args[0], args[1])], r, c)
        if fn == "SUMIFS":
            pairs = [(args[i], args[i + 1]) for i in range(1, len(args) - 1, 2)]
            return self.sumifs(args[0], pairs, r, c)
        return NAME

    def sumifs(self, total, pairs, r, c):
        """SUMIF/SUMIFS over a sheet: the rows whose every criterion holds."""
        if not isinstance(total, Rng):
            return VALUE
        tests = []
        for rng, criterion in pairs:
            if not isinstance(rng, Rng):
                return VALUE
            crit = self.get(criterion, r, c)
            if isinstance(crit, Err):
                return crit
            tests.append((rng, _criterion(crit)))
        grid = self._grid_of(total.sheet)
        last = min(total.r2, len(grid) - 1)
        out = 0.0
        for offset in range(0, last - total.r1 + 1):
            ok = True
            for rng, test in tests:
                v = self._grid_value(rng.sheet, rng.r1 + offset, rng.c1)
                if not test(v):
                    ok = False
                    break
            if not ok:
                continue
            v = self._grid_value(total.sheet, total.r1 + offset, total.c1)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                out += float(v)
        return out

    def indirect(self, ref):
        """``[2026_BO_Actual_Template.xlsx]KYC!$c:$R`` → a range on the
        template; ``1:8`` → whole rows here; anything else on this
        workbook."""
        if isinstance(ref, Err):
            return ref
        rows = re.match(r"^\s*\$?(\d+):\$?(\d+)\s*$", ref)
        if rows:
            return Rng(None, int(rows.group(1)) - 1, 0,
                       int(rows.group(2)) - 1, 16383)
        m = re.match(r"^\s*(?:\[[^\]]*\])?'?([^'!]+)'?!\$?([A-Za-z]+)\$?(\d*)"
                     r"(?::\$?([A-Za-z]+)\$?(\d*))?\s*$", ref)
        if not m:
            return REF
        sheet, c1, r1, c2, r2 = m.groups()
        external = ref.strip().startswith("[")
        col1 = _column_index(c1)
        col2 = _column_index(c2 or c1)
        row1 = int(r1) - 1 if r1 else 0
        row2 = int(r2) - 1 if r2 else 1048575
        if external:
            # Excel matches sheet names without case: the Commercial card
            # asks for "Trade_Finance_Income" and the template's sheet is
            # "Trade_Finance_income".
            sheet = _sheet_named(self.actuals, sheet)
            if sheet is None:
                return REF
            return Rng("@" + sheet, row1, col1, row2, col2)
        return Rng(_sheet_named(self.book["tables"], sheet) or sheet, row1,
                   col1, row2, col2)

    def vlookup(self, args, r, c):
        key = self.get(args[0], r, c)
        if isinstance(key, Err):
            return key
        table = args[1]
        if not isinstance(table, Rng):
            return table if isinstance(table, Err) else VALUE
        index = self.number(args[2], r, c)
        if isinstance(index, Err):
            return index
        index = int(index)
        exact = True
        if len(args) > 3:
            mode = self.get(args[3], r, c)
            exact = not _truth(mode) if not isinstance(mode, Err) else True
        if index < 1 or index > table.c2 - table.c1 + 1:
            return REF
        grid = self._grid_of(table.sheet)
        last = min(table.r2, len(grid) - 1)
        if exact:
            for row in range(table.r1, last + 1):
                if _lookup_equal(key, _plain(_at(grid, row, table.c1))):
                    found = self._grid_value(table.sheet, row,
                                             table.c1 + index - 1)
                    return 0.0 if found is None else found
            return NA
        best = None
        for row in range(table.r1, last + 1):
            v = _plain(_at(grid, row, table.c1))
            if v is None or isinstance(v, Err):
                continue
            if _compare("<=", v, key) is True:
                best = row
            else:
                break
        if best is None:
            return NA
        found = self._grid_value(table.sheet, best, table.c1 + index - 1)
        return 0.0 if found is None else found

    def _grid_of(self, sheet):
        if sheet and sheet.startswith("@"):
            return self.actuals.get(sheet[1:], [])
        if sheet in (None, self.sheet):
            return self.card["grid"]
        return self.book["tables"].get(sheet, [])

    def _grid_value(self, sheet, r, c):
        if sheet and sheet.startswith("@"):
            return _plain(_at(self.actuals.get(sheet[1:], []), r, c))
        return self.value(sheet, r, c)


def _at(grid, r, c):
    if 0 <= r < len(grid):
        row = grid[r]
        if 0 <= c < len(row):
            return row[c]
    return None


def _column_index(letters):
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _truth(v):
    if isinstance(v, Err):
        return v
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v != 0
    s = str(v).strip().upper()
    if s == "TRUE":
        return True
    if s == "FALSE":
        return False
    return VALUE


def _compare(op, a, b):
    """Excel's comparison: blank is 0 against a number and "" against text;
    text sorts after every number; text compares without case."""
    if a is None:
        a = "" if isinstance(b, str) else 0.0
    if b is None:
        b = "" if isinstance(a, str) else 0.0
    if isinstance(a, bool):
        a_rank, a = 2, float(a)
    elif isinstance(a, (int, float)):
        a_rank, a = 0, float(a)
    else:
        a_rank, a = 1, str(a).lower()
    if isinstance(b, bool):
        b_rank, b = 2, float(b)
    elif isinstance(b, (int, float)):
        b_rank, b = 0, float(b)
    else:
        b_rank, b = 1, str(b).lower()
    left, right = (a_rank, a), (b_rank, b)
    if a_rank != b_rank:
        left, right = (a_rank,), (b_rank,)
    return {"<": left < right, "<=": left <= right, "=": left == right,
            ">=": left >= right, ">": left > right,
            "<>": left != right}[op]


def _lookup_equal(key, value):
    if value is None or isinstance(value, Err):
        return False
    if isinstance(key, str) or isinstance(value, str):
        if not (isinstance(key, str) and isinstance(value, str)):
            return False
        # Exact match ignores case and nothing else: "Bob Waigwa " is not
        # "Bob Waigwa", in Excel or here.
        return key.lower() == value.lower()
    try:
        return float(key) == float(value)
    except (TypeError, ValueError):
        return False


def _sheet_named(sheets, name):
    """The key in ``sheets`` Excel would take ``name`` to mean, or None."""
    if name in sheets:
        return name
    lowered = name.lower()
    return next((key for key in sheets if key.lower() == lowered), None)


def _criterion(crit):
    """A SUMIF criterion → a test. ``"Bob"`` is equal-ignoring-case,
    ``">5"`` and friends compare, a number is numeric equality."""
    if isinstance(crit, str):
        m = re.match(r"^(<=|>=|<>|<|>|=)(.*)$", crit)
        if m:
            op, rest = m.groups()
            try:
                target = float(rest)
            except ValueError:
                target = rest
            return lambda v: (v is not None and not isinstance(v, Err)
                              and _compare(op, v, target) is True)
        return lambda v: isinstance(v, str) and v.lower() == crit.lower()
    if crit is None:
        return lambda v: v in (None, "")
    return lambda v: (isinstance(v, (int, float)) and not isinstance(v, bool)
                      and float(v) == float(crit))


def _round(fn, v, digits):
    factor = 10 ** digits
    x = v * factor
    if fn == "ROUND":
        y = math.floor(abs(x) + 0.5)
    elif fn == "ROUNDUP":
        y = math.ceil(abs(x) - 1e-9)
    else:
        y = math.floor(abs(x) + 1e-9)
    return math.copysign(y, x) / factor


def _general(v):
    if v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return repr(v)


def _format(v, fmt):
    if isinstance(v, Err):
        return v
    if isinstance(fmt, Err):
        return fmt
    if v is None:
        v = 0.0
    if not isinstance(v, (int, float)):
        return str(v)
    m = re.match(r"^0(?:\.(0+))?(%?)$", fmt or "")
    if not m:
        return _general(float(v))
    decimals = len(m.group(1) or "")
    if m.group(2):
        return f"{v * 100:.{decimals}f}%"
    return f"{v:.{decimals}f}"


_EPOCH = datetime.date(1899, 12, 30)


def _from_serial(serial):
    return _EPOCH + datetime.timedelta(days=int(serial))


def _to_serial(year, month, day):
    year += (month - 1) // 12
    month = (month - 1) % 12 + 1
    first = datetime.date(year, month, 1)
    return float((first - _EPOCH).days + day - 1)


# ─────────────────────────────────────────────────────────────────────────
# Cards
# ─────────────────────────────────────────────────────────────────────────

class Holder:
    """A List row: who, and which card."""

    def __init__(self, spec, row_index, row):
        def at(field):
            index = spec.list_cols.get(field)
            if index is None:
                return None
            return row[index] if index < len(row) else None

        self.book = spec.book
        self.row = row_index
        pf = at("pf")
        self.pf = (str(int(pf)) if isinstance(pf, float) and pf == int(pf)
                   else str(pf or "").strip())
        self.sales_code = str(at("code") or "").strip().upper()
        name = at("name")
        self.name = name if isinstance(name, str) else str(name or "")
        self.role = str(at("role") or "").strip()
        self.branch = at("branch")
        self.bbm = at("bbm")
        self.sheet = (spec.role_sheets.get(self.role)
                      or getattr(spec, "default_sheet", None))


def spec_of(book):
    return SPECS[book.get("book", BACK_OFFICE)]


def holders(book):
    spec = spec_of(book)
    grid = book["tables"].get("List", [])
    out = []
    for index, row in enumerate(grid[1:], start=1):
        holder = Holder(spec, index, row)
        if holder.name and (holder.sales_code or holder.pf):
            out.append(holder)
    return out


def _layout(spec, card):
    """Where the lines are: the header row says, and the end marker (the
    TOTAL row, or the signature block) ends them."""
    grid = card["grid"]
    header = next((i for i, row in enumerate(grid)
                   if any(isinstance(v, str) and v.strip() == spec.header
                          for v in row)), 7)
    end = next((i for i, row in enumerate(grid)
                if i > header and any(
                    isinstance(v, str) and v.strip().upper().startswith(
                        tuple(m.upper() for m in spec.ends))
                    for v in row)), len(grid))
    return header, end


def compute(book, actuals, holder, month, overrides=None):
    """One person's card for one month → dict, as the workbook would show it.

    ``overrides`` puts a value straight into a card cell - {(row, col):
    value} - which is how the live card hands the formulas today's figures
    in place of the template's.
    """
    spec = spec_of(book)
    card = book["cards"][holder.sheet]
    ev = Evaluator(book, actuals, holder, month)
    ev.inputs.update(overrides or {})
    col = spec.cols
    header, end = _layout(spec, card)

    def at(r, field):
        index = col.get(field)
        return None if index is None else _out(ev.cell(r, index))

    lines, perspective = [], ""
    for r in range(header + 1, end):
        weight = ev.cell(r, col["weight"])
        measure = ev.cell(r, col["measure"])
        initiative = ev.cell(r, col["initiative"])
        label = ev.cell(r, col["perspective"])
        if isinstance(label, str) and label.strip():
            perspective = " ".join(label.split())
        if not isinstance(weight, (int, float)) or isinstance(weight, bool):
            continue
        if not (_txt(measure) or _txt(initiative)):
            continue  # the weights' own total row
        feed = ev.cell(r, col["feed"])
        lines.append({
            "row": r + 1,
            "perspective": perspective,
            "measure": _txt(measure),
            "initiative": _txt(initiative),
            "weight": float(weight),
            "prior": at(r, "prior"),
            "growth": at(r, "growth"),
            "pm": at(r, "fy"),
            "ytd_target": at(r, "ytd"),
            "actual": at(r, "actual"),
            "score": at(r, "score"),
            "weighted": at(r, "weighted"),
            "feed": _txt(feed),
            # A line whose actual the card settles itself - typed
            # (complaints) or worked out on the card (compliments = its own
            # target) - rather than reading the template through INDIRECT.
            "typed": not any(
                t[0] == "f" and t[1] == "INDIRECT"
                for t in card["f"].get(f"{r},{col['actual']}", [])),
            "empty": (_feed_empty(actuals, feed, holder.name, month)
                      if spec.book == BACK_OFFICE else None),
        })
    score = ev.cell(*spec.score)
    return {"sheet": holder.sheet, "month": month, "lines": lines,
            "role": holder.role, "total": _out(score)}


def _txt(v):
    return " ".join(v.split()) if isinstance(v, str) else ""


def _out(v):
    if isinstance(v, Err):
        return {"e": v.code}
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return {"e": "#NUM!"}
    return v


def _feed_empty(actuals, feed, name, month):
    """Is this person's cell for the month blank on the template sheet?"""
    if not isinstance(feed, str) or feed not in actuals:
        return None
    for row in actuals[feed][1:]:
        if len(row) > 2 and isinstance(row[2], str) and isinstance(name, str) \
                and row[2].lower() == name.lower():
            column = 6 + month - 1
            value = row[column] if column < len(row) else None
            return value in (None, "")
    return None


def recorded_scores(book):
    """What the macro wrote, by person → {(sales code, month): score}.

    Matched on the sales code column, not the row: the List has been
    re-ordered since some of these months were written.
    """
    spec = spec_of(book)
    out = {}
    for sheet, header, code_col, first in spec.summaries:
        grid = book["tables"].get(sheet, [])
        for row in grid[header + 1:]:
            code = str(row[code_col] if len(row) > code_col else ""
                       ).strip().upper()
            if not code:
                continue
            for month in range(1, 13):
                column = first + month - 1
                v = row[column] if column < len(row) else None
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    out[(code, month)] = float(v)
    return out


# ─────────────────────────────────────────────────────────────────────────
# Storing, and showing
# ─────────────────────────────────────────────────────────────────────────

def _first_of(year, month):
    return datetime.date(year, month, 1)


def book_year(book):
    grid = book["tables"].get("Constants", [])
    year = _at(grid, *spec_of(book).year)
    if isinstance(year, (int, float)) and 2000 < year < 2100:
        return int(year)
    return datetime.date.today().year


def compute_all(book, actuals, months=None):
    """Every holder × every month with figures → [(holder, month, card)]."""
    months = months or template_months(actuals)
    out = []
    for holder in holders(book):
        if holder.sheet is None or holder.sheet not in book["cards"]:
            continue
        for month in months:
            out.append((holder, month, compute(book, actuals, holder, month)))
    return out


def store(results, year, recorded=None, loaded_by=""):
    """One row per person per month, replaced on reload → how many."""
    from django.db import transaction

    from .models import BoScorecard

    recorded = recorded or {}
    with transaction.atomic():
        for holder, month, card in results:
            total = card["total"]
            BoScorecard.objects.update_or_create(
                book=holder.book, sales_code=holder.sales_code or holder.pf,
                month=_first_of(year, month),
                defaults={
                    "staff_pf": holder.pf, "staff_name": holder.name,
                    "sheet": holder.sheet, "branch": str(holder.branch or ""),
                    "card": card,
                    "total": total if _num(total) else None,
                    "recorded": recorded.get((holder.sales_code, month)),
                    "loaded_by": loaded_by,
                })
    return len(results)


def save_file(kind, payload, file_name="", loaded_by=""):
    from .models import BoWorkbookFile

    BoWorkbookFile.objects.update_or_create(
        kind=kind, defaults={"payload": payload, "file_name": file_name,
                             "loaded_by": loaded_by})


def saved(kind):
    from .models import BoWorkbookFile

    row = BoWorkbookFile.objects.filter(kind=kind).first()
    return row.payload if row is not None else None


def check(book, actuals):
    """What loading would do → a report. Nothing is written.

    The proof is the macro's own record: every month it wrote on
    Monthly_Summ is recomputed and compared, by sales code.
    """
    spec = spec_of(book)
    months = template_months(actuals)
    people = [h for h in holders(book) if h.sheet in book["cards"]]
    recorded = recorded_scores(book)
    by_month, differ = {}, []
    for holder in people:
        for month in months:
            score = recorded.get((holder.sales_code, month))
            if score is None:
                continue
            total = compute(book, actuals, holder, month)["total"]
            tally = by_month.setdefault(month, {"recorded": 0, "same": 0})
            tally["recorded"] += 1
            if _num(total) and abs(total - score) < 1e-9:
                tally["same"] += 1
            else:
                differ.append({"sales_code": holder.sales_code,
                               "name": holder.name, "role": holder.role,
                               "month": month, "recorded": score,
                               "computed": total})
    roles = {}
    for holder in holders(book):
        roles[holder.role or "(blank)"] = roles.get(holder.role or "(blank)",
                                                    0) + 1
    return {
        "book": spec.book,
        "year": book_year(book),
        "months": months,
        "people": len(people),
        "cards": sorted({h.sheet for h in people}),
        "roles_without_a_card": sorted(
            r for r in roles
            if r not in spec.role_sheets and r != "(blank)"),
        "left": sum(1 for h in people
                    if str(h.branch or "").strip().lower() == "left"),
        "against_the_macro": [
            {"month": m, "recorded": t["recorded"], "same": t["same"]}
            for m, t in sorted(by_month.items())],
        "differences": differ[:200],
        "problems": book.get("problems", []),
    }


def rebuild(book, actuals, loaded_by=""):
    """Compute and store every card for every month → how many stored."""
    results = compute_all(book, actuals)
    return store(results, book_year(book), recorded_scores(book),
                 loaded_by=loaded_by)


def stored_for(sales_code, as_at=None, book=BACK_OFFICE):
    from .models import BoScorecard

    code = (sales_code or "").strip()
    rows = BoScorecard.objects.filter(book=book, sales_code__iexact=code)
    if as_at is not None:
        rows = rows.filter(month__lte=as_at)
    return rows.order_by("-month").first()


def stored_month(sales_code, month_start, book):
    """The stored card for exactly this month, or None."""
    from .models import BoScorecard

    return BoScorecard.objects.filter(
        book=book, sales_code__iexact=(sales_code or "").strip(),
        month=month_start).first()


def months_loaded(sales_code, book=BACK_OFFICE):
    from .models import BoScorecard

    return [m.strftime("%Y-%m") for m in BoScorecard.objects.filter(
        book=book, sales_code__iexact=(sales_code or "").strip())
        .order_by("month").values_list("month", flat=True)]


def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _month_end(first):
    nxt = (first.replace(day=28) + datetime.timedelta(days=4)).replace(day=1)
    return nxt - datetime.timedelta(days=1)


def card_payload(record, asked=None):
    """A stored workbook card in the shape every card has."""
    spec = SPECS.get(getattr(record, "book", BACK_OFFICE) or BACK_OFFICE,
                     SPECS[BACK_OFFICE])
    data = record.card or {}
    perspectives, seen = [], {}
    scored = unscored = 0.0
    month_end = _month_end(record.month)
    for line in data.get("lines", []):
        name = line["perspective"] or "Other"
        if name not in seen:
            seen[name] = {"perspective": name, "weight": 0.0, "lines": []}
            perspectives.append(seen[name])
        seen[name]["weight"] += line["weight"]
        score, weighted = line.get("score"), line.get("weighted")
        if _num(score):
            scored += line["weight"]
        else:
            unscored += line["weight"]
        pm, ytd, actual = line.get("pm"), line.get("ytd_target"), \
            line.get("actual")
        prior, growth = line.get("prior"), line.get("growth")
        error = next((v["e"] for v in (actual, score) if isinstance(v, dict)),
                     "")
        pending = ""
        if error:
            pending = (f"The workbook itself shows {error} here: "
                       f"{record.staff_name} has no row on the "
                       f"{line.get('feed') or 'actuals'} sheet of the actuals "
                       f"template." if error == "#N/A" else
                       f"The workbook itself shows {error} on this line.")
        as_at = f"as at {month_end.day} {month_end:%b %Y}"
        if line.get("typed"):
            as_at = "typed on the card in the workbook"
        elif line.get("empty"):
            as_at = (f"no figure for {month_end:%B} on the "
                     f"{line.get('feed')} sheet - read as 0, as the workbook "
                     f"reads it")
        seen[name]["lines"].append({
            "kpi_code": (line.get("feed") or line["measure"]).lower(),
            "kpi_name": line["measure"] or line["initiative"],
            "measure_of_success": line["initiative"] or line["measure"],
            "weight": line["weight"],
            "ytd_target": ytd if _num(ytd) else None,
            "annual_target": pm if _num(pm) else None,
            "target_text": pm if isinstance(pm, str) else "",
            "prior_year": prior if _num(prior) else None,
            "growth_target": growth if _num(growth) else None,
            "target_field": "",
            "target_source": spec.source,
            "full_year_target": pm == ytd,
            "ytd_actual": actual if _num(actual) else None,
            "warehouse_said": None,
            "score": score if _num(score) else None,
            "weighted_score": weighted if _num(weighted) else None,
            "counted_in_total": True,
            "as_at": as_at,
            "as_at_date": month_end.isoformat(),
            "higher_is_better": True,
            "full_marks_rule": "",
            "score_cap": None,
            "pending": pending,
            "pending_label": "Workbook error" if pending else "",
            "diagnostic": "",
        })
    total = record.total
    sheet = record.sheet
    empty = sum(1 for ln in data.get("lines", [])
                if ln.get("empty") and not ln.get("typed"))
    # The macro's own record wins: it is what the desk sent, with the List
    # and the targets as they were that month. The workbook carries only
    # today's - Q2's revised targets, today's branches, leavers marked
    # "left" - so recomputing an old month can come out differently, and
    # where it does the card says so rather than quietly replacing a score
    # somebody signed.
    recorded = record.recorded
    notes = []
    if _num(recorded) and not (_num(total) and abs(recorded - total) < 1e-9):
        notes.append(
            f"The desk's scorecard run recorded {recorded:.1%} for "
            f"{record.month:%B %Y}, with the targets and List as they were "
            f"then, and that is the score shown. The lines below are the "
            f"month recomputed from the workbook as it is now (revised "
            f"targets, current branches), which gives "
            + (f"{total:.1%}." if _num(total) else
               "an error in the workbook itself."))
        total = recorded
    if empty:
        notes.append(
            f"{empty} line(s) had no figure for {record.month:%B %Y} on the "
            f"actuals template, so they read as 0 - as the workbook reads "
            f"them.")
    return {
        "has_card": True,
        "live": False,
        "workbook_card": spec.book,
        "back_office_workbook": spec.book == BACK_OFFICE,
        "staff": {
            "sales_code": record.sales_code,
            "name": record.staff_name,
            "title": (TITLES.get(sheet, sheet) if spec.book == BACK_OFFICE
                      else getattr(spec, "title", "")
                      or data.get("role") or sheet),
            "role_code": (ROLE_CODES.get(sheet, sheet)
                          if spec.book == BACK_OFFICE else
                          getattr(spec, "role_code", "")
                          or SALES_ROLE_CODES.get(data.get("role", ""),
                                                  sheet)),
            "branch": record.branch,
            "team_leader": "",
            "roster": f"{spec.source}'s List",
            "started": None,
        },
        "book": {"customers": None, "note": ""},
        "period": {"label": record.month.strftime("%B %Y"),
                   "as_at": month_end.isoformat()},
        "target_source": spec.source,
        "performance_score": round(total, 4) if _num(total) else None,
        "scored_lines": sum(1 for p in perspectives for ln in p["lines"]
                            if ln["score"] is not None),
        "pending_lines": sum(1 for p in perspectives for ln in p["lines"]
                             if ln["pending"]),
        "scored_weight": round(scored, 4),
        "unscored_weight": round(unscored, 4),
        "perspectives": perspectives,
        "months_loaded": months_loaded(record.sales_code, spec.book),
        "coverage_note": " ".join(notes),
        "showing_other_month": (
            "" if asked is None or asked.replace(day=1) == record.month else
            f"Nothing is loaded for {asked:%B %Y}; this is "
            f"{record.month:%B %Y}, the latest month on or before it."),
        "how_scored": (f"Computed from {spec.source}'s own formulas over "
                       f"the actuals template - the same cells the desk's "
                       f"macro fills and reads, so the score is the one the "
                       f"workbook gives."),
    }


# ─────────────────────────────────────────────────────────────────────────
# The live card: the workbook's rules, today's figures
# ─────────────────────────────────────────────────────────────────────────

def _key(text):
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def waiting_on(sales_code):
    """A loaded workbook's List holds this person but no card of theirs is
    computed yet → the sentence saying which file is missing, or None.

    The cards are computed only once both a workbook and the actuals
    template it reads are loaded; until then the person is on no roster that
    gives them a card, and "no scorecard for your role" would be wrong.
    """
    code = (sales_code or "").strip().upper()
    if not code:
        return None
    for spec in SPECS.values():
        book = saved(spec.files[0])
        if book is None or not any(h.sales_code == code
                                   for h in holders(book)):
            continue
        if saved(spec.files[1]) is None:
            return (f"You are on the {spec.label} scorecards workbook, but "
                    f"the actuals template it reads has not been loaded, so "
                    f"no card is computed yet. Administration loads it under "
                    f"Scorecards → Figures to load → Load the scorecard "
                    f"workbooks.")
        return (f"You are on the {spec.label} scorecards workbook, but no "
                f"card was computed for you - load the actuals template "
                f"again under Scorecards → Figures to load.")
    return None


def _put_template(actuals, sheet, code, month, value, block=0):
    """Write one person's figure for one month into a template copy → True
    when their row was there. ``block`` is which of a sheet's side-by-side
    tables (20 columns apart) - the bancassurance sheet keeps three."""
    from .bb_live_feeds import BLOCK_WIDTH

    o = block * BLOCK_WIDTH
    name = _sheet_named(actuals, sheet)
    for row in (actuals.get(name) or [] if name else [])[1:]:
        if len(row) > o + 1 and                 str(row[o + 1] or "").strip().upper() == code.upper():
            while len(row) < o + 6 + month:
                row.append(None)
            row[o + 6 + month - 1] = value
            return True
    return False


#: Template sheets that only grow through the year: counts and sums of what
#: was sold, opened or disbursed. Balances, rates and charges are not here.
ACCUMULATING_SHEETS = {"plotsalesvolume", "plotsalesvalue",
                       "tradefinanceincome", "bancaassuranceincome",
                       "newcustomers", "assetdrawdown", "traininghours"}


def _get_template(actuals, sheet, code, month, block=0):
    """One person's figure for one month on the template, or None."""
    from .bb_live_feeds import BLOCK_WIDTH

    o = block * BLOCK_WIDTH
    name = _sheet_named(actuals, sheet)
    for row in (actuals.get(name) or [] if name else [])[1:]:
        if len(row) > o + 1 and \
                str(row[o + 1] or "").strip().upper() == code.upper():
            return row[o + 5 + month] if len(row) > o + 5 + month else None
    return None


def _carry_forward(actuals, sheet, code, month):
    """Put the desk's latest figure for this person into ``month`` on a
    template copy → the month it came from, or None.

    The template is year to date (balances, YTD sums, rates), so the desk's
    latest month IS the person's year to date - a line the warehouse has no
    live figure for shows that and is scored, rather than reading nought
    while the person may already have made the target. Every side-by-side
    table of the sheet is carried, so each line finds its own. "Latest" is
    the last month with any figure for anyone on that table: a formula's
    nought in a month nobody has filled is not the desk's figure.
    """
    from .bb_live_feeds import BLOCK_WIDTH

    name = _sheet_named(actuals, sheet)
    rows = (actuals.get(name) or []) if name else []
    if len(rows) < 2:
        return None
    width = max(len(r) for r in rows)
    used = None
    for o in range(0, width, BLOCK_WIDTH):
        filled = [r for r in rows[1:] if len(r) > o + 1 and r[o + 1]]
        mine = next((r for r in filled
                     if str(r[o + 1]).strip().upper() == code.upper()), None)
        if mine is None:
            continue
        latest = next((m for m in range(month, 0, -1)
                       if any(len(r) > o + 5 + m and _num(r[o + 5 + m])
                              and r[o + 5 + m] != 0 for r in filled)), None)
        if latest is None:
            continue
        if latest < month:
            while len(mine) < o + 6 + month:
                mine.append(None)
            mine[o + 5 + month] = (mine[o + 5 + latest]
                                   if len(mine) > o + 5 + latest else None)
        used = max(used or 0, latest)
    return used


def _line_blocks(book, actuals, holder, month, base, feeds_live):
    """{card row: block} for the lines on a sheet with more than one table
    proven: a marker is put in each block in turn, and the lines whose
    actual moves are the ones reading it. Exact, where guessing the block
    from a line's label would not be."""
    import copy

    out = {}
    for wl in base["lines"]:
        blocks = feeds_live.get(_key(wl["feed"])) or {}
        if set(blocks) == {0} or wl["row"] in out:
            continue
        for b in blocks:
            trial = copy.deepcopy(actuals)
            if not _put_template(trial, wl["feed"], holder.sales_code, month,
                                 987654321.0, b):
                continue
            moved = compute(book, trial, holder, month)
            for before, after in zip(base["lines"], moved["lines"]):
                if before["actual"] != after["actual"]:
                    out[after["row"]] = b
    return out


def live_card(sales_code, engine_card, as_at):
    """The month still open, scored by the workbook's own formulas over the
    figures the warehouse gives today → a card, or None.

    A closed month is the stored card - the template the desk filled, the
    score the desk sent. The open month has no template yet, so each line's
    actual is the live engine's warehouse figure for that line, put into the
    card's actual cell; the workbook's own targets, caps and weights then
    score it. A line the warehouse has no figure for is shown as waiting for
    one, never as a zero.
    """
    from .models import BoScorecard
    from .scorecard_automation.models import ScKpi

    code = (sales_code or "").strip()
    latest = (BoScorecard.objects.filter(sales_code__iexact=code)
              .order_by("-month").first())
    if latest is None:
        return None
    spec = SPECS.get(latest.book, SPECS[BACK_OFFICE])
    book, actuals = saved(spec.files[0]), saved(spec.files[1])
    if book is None or actuals is None:
        return None
    holder = next((h for h in holders(book)
                   if h.sales_code == code.upper() and h.sheet), None)
    if holder is None or holder.sheet not in book["cards"]:
        return None

    if as_at is None:
        first_this = datetime.date.today().replace(day=1)
        as_at = first_this - datetime.timedelta(days=1)
    if as_at.year != book_year(book):
        return None
    month = as_at.month

    # The live engine's figure for each of its lines, labelled with the
    # template sheet that line reads - which is what the workbook names too.
    live_lines = [ln for p in (engine_card or {}).get("perspectives", [])
                  for ln in p["lines"]]
    feeds = dict(ScKpi.objects.filter(
        kpi_code__in=[ln["kpi_code"] for ln in live_lines])
        .values_list("kpi_code", "actuals_sheet"))
    pool = [{"feed": _key(feeds.get(ln["kpi_code"]) or ""),
             "name": _key(ln.get("kpi_name")), "line": ln}
            for ln in live_lines]

    # Pass one with the template's own figures for the month, so the rows
    # and feeds are known; then the same card with today's figures in.
    base = compute(book, actuals, holder, month)
    actual_col = spec.cols["actual"]
    overrides, found = {}, {}
    # The back office first reads the feeds the probe has proven - a count
    # off transaction_diary that gave the desk's own figures for a closed
    # month (bo_live_feeds). They win over the engine, whose readers look
    # up a seller code that a teller does not sell under.
    # A branch manager's row is the branch's own total, so the branch
    # feeds the probe has proven (bb_live_feeds) are read by branch - the
    # engine's per-person readers never give a branch total.
    feeds_live, reader = {}, None
    if spec.book == BACK_OFFICE:
        from . import bo_live_feeds

        feeds_live = {_key(k): v for k, v in bo_live_feeds.proven().items()}
        reader = lambda feed: bo_live_feeds.live_value(feed, code, as_at)  # noqa: E731
    elif spec.book == BRANCH_BUSINESS:
        from . import bb_live_feeds

        # sheet → {block: feed}: one sheet can hold several proven tables.
        for k, v in bb_live_feeds.proven().items():
            feeds_live.setdefault(_key(k.split("#")[0]), {})[
                int(v.get("block") or 0)] = v
        reader = lambda feed: bb_live_feeds.live_value(  # noqa: E731
            feed, holder.branch, as_at)
        blocks = _line_blocks(book, actuals, holder, month, base, feeds_live)
    live_rows, template_live = set(), None
    # Every line starts from the desk's latest year-to-date figure; a live
    # one then writes over it. The back office template keys its rows by
    # name on some sheets, so it is not carried this way.
    carried = {}
    if spec.book != BACK_OFFICE:
        import copy

        template_live = copy.deepcopy(actuals)
        by_sheet = {}
        for wl in base["lines"]:
            if wl["typed"] or not wl["feed"]:
                continue
            k = _key(wl["feed"])
            if k not in by_sheet:
                by_sheet[k] = _carry_forward(template_live, wl["feed"],
                                             code, month)
            if by_sheet[k]:
                carried[wl["row"]] = by_sheet[k]
    # The desk's year to date per line, before anything live is put in.
    desk = ({wl["row"]: wl["actual"] for wl in
             compute(book, template_live, holder, month)["lines"]}
            if carried else {})
    prior = {}
    for wl in base["lines"]:
        if wl["typed"]:
            continue
        feed = feeds_live.get(_key(wl["feed"]))
        if feed is not None and spec.book == BRANCH_BUSINESS:
            only_first = set(feed) == {0}
            block = 0 if only_first else blocks.get(wl["row"])
            feed = feed.get(block) if block is not None else None
        if feed is not None:
            value = reader(feed)
            if value is not None and spec.book == BRANCH_BUSINESS:
                # Into the template's cell for the month, not the card's:
                # the card's own formula takes the 2025 base off a balance
                # (deposits, loans), and writing over that cell would lose it.
                if template_live is None:
                    import copy

                    template_live = copy.deepcopy(actuals)
                prior[wl["row"]] = (wl["feed"], block, _get_template(
                    template_live, wl["feed"], code, month, block))
                if _put_template(template_live, wl["feed"], code, month,
                                 float(value), block):
                    live_rows.add(wl["row"])
                continue
            if value is not None:
                overrides[(wl["row"] - 1, actual_col)] = float(value)
                continue
        if spec.book == BRANCH_BUSINESS:
            continue
        match = (next((p for p in pool if p["feed"] and
                       p["feed"] == _key(wl["feed"])), None)
                 or next((p for p in pool if p["name"] and
                          p["name"] == _key(wl["measure"])), None))
        if match is None:
            continue
        pool.remove(match)
        value = match["line"].get("ytd_actual")
        found[wl["row"]] = match["line"]
        if _num(value):
            overrides[(wl["row"] - 1, actual_col)] = float(value)
    card = compute(book, template_live or actuals, holder, month, overrides)

    # A year-to-date count or sum cannot fall below what the desk already
    # recorded for the year. Where a live reader gives less - property put
    # on nobody because the seller's name is spelt differently, a register
    # not yet reloaded - the desk's figure stands.
    dropped = []
    for wl in card["lines"]:
        r = wl["row"]
        is_live = (r - 1, actual_col) in overrides or r in live_rows
        if (is_live and r in carried
                and _key(wl["feed"]) in ACCUMULATING_SHEETS
                and _num(wl["actual"]) and _num(desk.get(r))
                and wl["actual"] < desk[r]):
            dropped.append(r)
            overrides.pop((r - 1, actual_col), None)
            if r in live_rows:
                live_rows.discard(r)
                sheet, block, old = prior[r]
                _put_template(template_live, sheet, code, month, old, block)
    if dropped:
        card = compute(book, template_live or actuals, holder, month,
                       overrides)

    total, waiting = 0.0, 0
    for wl in card["lines"]:
        live = ((wl["row"] - 1, actual_col) in overrides
                or wl["row"] in live_rows)
        if wl["typed"] or live:
            if _num(wl["weighted"]):
                total += wl["weighted"]
            continue
        if wl["row"] in carried and _num(wl["actual"]):
            wl["carried"] = carried[wl["row"]]
            if _num(wl["weighted"]):
                total += wl["weighted"]
            continue
        waiting += 1
        engine = found.get(wl["row"]) or {}
        wl["actual"] = wl["score"] = wl["weighted"] = None
        wl["waiting"] = (engine.get("pending") or
                         "The warehouse has no figure for this line yet.")
    card["total"] = total

    record = type("LiveRecord", (), {})()
    record.book, record.card, record.total = spec.book, card, total
    record.recorded, record.month = None, as_at.replace(day=1)
    record.sales_code, record.staff_name = code.upper(), holder.name
    record.sheet, record.branch = holder.sheet, str(holder.branch or "")
    payload = card_payload(record)
    for p in payload["perspectives"]:
        for ln, wl in zip(p["lines"], [w for w in card["lines"]
                                       if (w["perspective"] or "Other")
                                       == p["perspective"]]):
            if wl.get("waiting"):
                ln["pending"] = wl["waiting"]
                ln["pending_label"] = "Waiting for data"
                ln["as_at"] = ""
            elif wl.get("carried"):
                # Every template sheet is year to date except turnaround
                # time, which is the month's own rate.
                kind = ("latest month" if "tat" in _key(wl["feed"])
                        else "year to date")
                ln["as_at"] = (
                    f"{kind}, the desk's latest figure "
                    f"({datetime.date(as_at.year, wl['carried'], 1):%b %Y})")
            else:
                ln["as_at"] = f"live, as at {as_at:%d %b %Y}"
    payload.update({
        "live": True,
        "pending_lines": waiting,
        "scored_lines": len(card["lines"]) - waiting,
        # Every month of the year to now: stored where the desk closed it,
        # live where it has not.
        "months_loaded": [f"{as_at.year}-{m:02d}"
                          for m in range(1, as_at.month + 1)],
        "coverage_note": (
            "" if not waiting else
            f"{waiting} line(s) have no live figure yet and are not counted; "
            f"they fill in as the warehouse has them."),
        "how_scored": (f"Live: today's warehouse figures, scored by "
                       f"{spec.source}'s own targets, caps and weights. "
                       f"Closed months show the card the desk sent."),
    })
    return payload
