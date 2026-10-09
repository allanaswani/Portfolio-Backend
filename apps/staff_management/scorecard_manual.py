"""The figures the system cannot compute, loaded by an administrator.

Most of a scorecard is now worked out from data the bank already holds. What is
left is measured somewhere this platform cannot reach — the customer survey
(NPS), the learning system (training hours), HR (leave), Internal Audit, the
branch's own engagement count, iApply covenant tracking, account-planning
reviews — plus any target the DMC load does not carry.

Those have nowhere to come from except a person typing them in, so this is the
screen where a person types them in. **Admin and superusers only**: this sets
what everybody is measured on.

Two tables, both of which already existed and neither of which anything was
writing:

* ``sc_employee_kpi_targets``            — a target for one person
* ``sc_employee_performance_actual_values`` — an actual for one person, one month

The card reads an uploaded figure only where it has nothing of its own. A
warehouse figure always wins, so loading a stale number here can never
overwrite a live one — and the card says which of the two it used, so a
hand-typed figure is never mistaken for a measured one.

The template is generated, not kept as a file
─────────────────────────────────────────────
It is built from the roster and the seeded cards at the moment it is asked
for: one row per person per line that actually needs a manual figure, with
their name, sales code, role and the KPI already filled in. A blank template
that somebody has to match up by hand is how a figure ends up against the
wrong sales code.
"""

from __future__ import annotations

import datetime
from io import BytesIO

#: KPI codes whose ACTUAL cannot be computed from anything the bank holds here,
#: with where the real number lives. The reason is shown on the template and on
#: the card, so nobody has to guess who to ask.
MANUAL_ACTUALS = {
    "nps": "Customer survey — Net Promoter Score",
    "portfolio_nps": "Customer survey — portfolio NPS",
    "portfolo_nps": "Customer survey — portfolio NPS",
    "portfolio_coverage_engagement": "Engagement count from the branch's own return",
    "business_banking_training": "Learning system — training hours completed",
    "personal_banking_training": "Learning system — training hours completed",
    "leave_management": "HR — leave days cleared",
    "audit": "Internal Audit — annual audit score",
    "errors": "Mortgage operations — application error rate",
    "weighted_sales": "Weighted sales dashboard",
    "weighted_sales_dashboard": "Weighted sales dashboard",
    "digital_adoption": "Digital channels — adoption count",
    "banking_covenant_tracking_should_be_in_iapply": "iApply — covenant tracking",
    "tooling_and_account_planning_all_customers": "Account planning reviews completed",
    "portfolio_management_aum_dec_previous_year": "Assets under management",
    "active_customers": "Customers active two months running",
    "ultimate_rm_new_customers": "New customers holding more than two products",
    "deposits_energy_water": "Energy & Water sector deposits",
    "net_mortgage_sales_10_commercial_rate_mortgages":
        "Mortgage sales at the market rate",
    "net_mortgage_sales_10_non_commercial_rate_mortgages":
        "Mortgage sales at the non-market rate",
    "bancassurance_premiums": "Bancassurance — renewals and other products",
    "other_bancassurance_premiums": "Bancassurance — other products",

    # ── the back-office cards ──
    # Operations staff are measured on controls and service rather than on a
    # balance: audit scores, turnaround, KYC cleanup, complaints logged,
    # operation losses, cash management. Almost none of that is in a warehouse
    # table, which is a statement about where the data lives and not a gap in
    # the card. Each names the return it comes from, so whoever has to find the
    # number knows who to ask.
    "bo_nps": "Customer survey — Net Promoter Score",
    "bo_crm": "CRM — complaints logged",
    "bo_kyc": "KYC cleanup return",
    "bo_branch_audit_branch_audit": "Internal Audit — branch audit score",
    "bo_branch_audit_resolution": "Internal Audit — repeat findings resolved",
    "bo_branch_audit_compliance": "Internal Audit — compliance score",
    "bo_branch_audit_cash_management": "Internal Audit — cash management",
    "bo_diaspora_audit": "Internal Audit — the diaspora return",
    "bo_cash_management": "Cash management return",
    "bo_operation_lossess": "Operational loss and fraud register",
    "bo_trx_errors": "Transaction error return",
    "bo_rbo_account_errors": "Back-office returns register",
    "bo_trx_productivity": "Teller transaction count",
    "bo_sales_productivity": "Team productivity return",
    "bo_training_hours": "Learning system — training hours completed",
    "bo_leave_management": "HR — leave days cleared",
    "bo_events": "Branch events return",
    "bo_branch_rtgs": "RTGS channel return",
    "bo_digital_activation": "Digital activation return",
    "bo_digital_customers": "Digital onboarding return",
    "bo_re_activation": "Dormancy reactivation return",
    "bo_active_customers": "Active-customer count",
    "bo_account_opening_pb": "Account-opening turnaround — personal banking",
    "bo_account_opening_bb": "Account-opening turnaround — business banking",
    "bo_account_opening_ub": "Account-opening turnaround — ultimate banking",

    # The sixteen whose own card names a feed the actuals template does not
    # have, or names none at all. Loading them by hand is the only route until
    # the desk says what they should read.
    "bo_account_errors": "Account errors — the card names a sheet the "
                         "actuals template does not have",
    "bo_account_opening_sme": "SME account-opening turnaround — no such "
                              "sheet on the actuals template",
    "bo_bbm_financials": "Branch financials — no such sheet on the actuals "
                         "template",
    "bo_deposit_growth": "Deposit growth — no deposits sheet on the "
                         "actuals template",
    "bo_cost": "Cost reduction — there is no operating-cost data anywhere "
               "in this warehouse",
    "bo_npl": "NPL — the card names a sheet the template does not have",
    "bo_rbo_monthly_summ": "Team productivity — the card names a summary "
                           "sheet rather than a feed",
    "bo_npl_reduce_p_l_provisions_from_diaspora":
        "Diaspora provisions specifically",
    "bo_customer_feedback_nil_client_complaints": "Client complaints",
    "bo_customer_feedback_achieve_5_client_compliments": "Client compliments",
    "bo_customer_feedback_achieve_3_ultimate_banking_client_co":
        "Ultimate Banking client compliments",
    "bo_customer_feedback_achieve_3_diaspora_banking_client_co":
        "Diaspora Banking client compliments",
    "bo_customer_journey_document_and_approve_the_customer_ser":
        "Whether the customer service journey is documented and approved",
    "bo_tds_track_ub_term_deposit_renewals":
        "Ultimate Banking term-deposit renewals tracked",
    "bo_term_deposits_track_diaspora_term_deposit_renewals_ret":
        "Diaspora term-deposit renewals tracked",
    "bo_tools_achieve_tat_of_1_day_for_mobile_banking_onboardi":
        "Mobile banking onboarding turnaround",
}

#: KPI codes with no target anywhere: the DMC load carries no column for them
#: and they are per person, so no role-wide figure can stand in.
MANUAL_TARGETS = {
    "direct_portfolio_contribution": "target_pbt_revenue is on the BRANCH DMC "
                                     "table only, and a branch revenue target "
                                     "is not one RM's target",
    "income_contribution": "as above",
    "income_contribution_portfolio_ntb": "as above",
    "operating_profit": "as above",
    "loan_loss": "target_loan_provisions is on the BRANCH DMC table only",
    "reduce_p_l_provisions_nil_provisions_loan_loss": "as above",
    "reduce_p_l_provisions_nill_provisions_loan_loss": "as above",
    "asset_growth": "no column on the per-person DMC table; the card currently "
                    "derives it as a rate on your own December book, and a "
                    "figure loaded here overrides that",
    # These three had a role target on the seeded cards, but it was one named
    # RM's own closing figure rather than a rate everybody shares, so 0029
    # cleared it. They are per person and there is no column for them.
    "digital_adoption": "per person, and no DMC column carries it",
    "portfolio_management_aum_dec_previous_year":
        "per person, and no DMC column carries it",
    "weighted_sales": "the weighted-sales dashboard target, per person",
}

COLUMNS = ["sales_code", "staff_name", "staff_role", "kpi_code", "kpi_name",
           "what_this_is", "target", "actual", "month"]


def _roster():
    """Everybody with a card, from all three rosters.

    The sales DMC tables answer for RMs. The back-office roster in
    ``employee_role_history`` answers for branch operations managers, customer
    service officers, tellers and the back office, who are on no sales roster -
    that table is keyed on sales targets a teller does not have.

    Listing only the sales side is how 101 people would have had figures that
    nobody could load.
    """
    from .live_scorecard import role_code_for
    from .models import (
        BranchEmployeeDmcData, BranchFinalEmployeeDmcData,
        EmployeeRoleHistory, StaffEmployeeData)

    seen, out = set(), []
    for model in (BranchEmployeeDmcData, BranchFinalEmployeeDmcData):
        rows = model.objects.exclude(sales_code__isnull=True).exclude(
            sales_code="").order_by("sales_code")
        for row in rows:
            code = (row.sales_code or "").strip().upper()
            if not code or code in seen:
                continue
            role = role_code_for(row.staff_role)
            if role is None:
                continue
            seen.add(code)
            out.append((code, row.staff_name or "", row.staff_role or "", role))

    # The back office. Names come from the HR record where there is one, and
    # from the roster's own notes otherwise - which is where migration 0032 put
    # them, because employee_role_history has nowhere else for a name.
    names = {(code or "").strip().upper(): name for code, name
             in StaffEmployeeData.objects.exclude(sales_code="")
             .values_list("sales_code", "staff_name")}
    for entry in (EmployeeRoleHistory.objects
                  .filter(role_code__startswith="bo_").order_by("sales_code")):
        code = (entry.sales_code or "").strip().upper()
        if not code or code in seen:
            continue
        seen.add(code)
        name = names.get(code) or (entry.notes or "").split("|")[0].strip()
        out.append((code, name, entry.role_code, entry.role_code))
    return out


def rows_needing_a_figure(sales_code=""):
    """One row per person per line that needs a figure typed in.

    Built from the roster and the seeded cards, so it is only ever lines that
    are actually on somebody's card - a template listing KPIs nobody is
    measured on is a template people stop reading.
    """
    from .scorecard_automation.models import ScKpi, ScRoleKpiMapping

    wanted = (sales_code or "").strip().upper()
    names = dict(ScKpi.objects.values_list("kpi_code", "kpi_name"))
    by_role = {}
    for mapping in ScRoleKpiMapping.objects.order_by("role_code", "kpi_order"):
        by_role.setdefault(mapping.role_code, []).append(mapping.kpi_code)

    out = []
    for code, name, role_title, role in _roster():
        if wanted and code != wanted:
            continue
        for kpi in by_role.get(role, []):
            needs_actual = kpi in MANUAL_ACTUALS
            needs_target = kpi in MANUAL_TARGETS
            if not (needs_actual or needs_target):
                continue
            out.append({
                "sales_code": code,
                "staff_name": name,
                "staff_role": role_title,
                "kpi_code": kpi,
                "kpi_name": names.get(kpi, kpi.replace("_", " ").title()),
                "what_this_is": (MANUAL_ACTUALS.get(kpi)
                                 or MANUAL_TARGETS.get(kpi, "")),
                "needs_actual": needs_actual,
                "needs_target": needs_target,
            })
    return out


def build_template(sales_code="", month=None):
    """The upload template, generated from the roster → ``BytesIO``.

    Pre-filled with who and which line, so the only thing to type is the
    number. Any existing figure is filled in too, which makes this the way to
    correct one as well as to add one.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from .scorecard_automation.models import (
        ScEmployeeKpiTarget, ScEmployeePerformanceActual)

    month = month or _default_month()
    year = month.year
    rows = rows_needing_a_figure(sales_code)

    targets = {(t.sales_code.strip().upper(), t.kpi_code): t.kpi_target
               for t in ScEmployeeKpiTarget.objects.filter(year=year)}
    actuals = {(a.sales_code or "").strip().upper() + "|" + a.kpi_code: a.kpi_value
               for a in ScEmployeePerformanceActual.objects.filter(eom_date=month)}

    wb = Workbook()
    ws = wb.active
    ws.title = "Figures"

    ws.cell(row=1, column=1, value=(
        "Type the number in the target or actual column and upload this file "
        "again. Leave a cell blank to leave that figure alone.")).font = Font(
            italic=True, color="FF8A6A2A")
    ws.cell(row=2, column=1, value=(
        "A figure loaded here is only used where the system has nothing of its "
        "own. A warehouse figure always wins.")).font = Font(
            italic=True, color="FF8A6A2A")

    header = 4
    for index, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header, column=index, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFF3F8F9")

    for offset, row in enumerate(rows, start=header + 1):
        key = (row["sales_code"], row["kpi_code"])
        values = [
            row["sales_code"], row["staff_name"], row["staff_role"],
            row["kpi_code"], row["kpi_name"], row["what_this_is"],
            targets.get(key) if row["needs_target"] else "n/a",
            actuals.get(f"{row['sales_code']}|{row['kpi_code']}")
            if row["needs_actual"] else "n/a",
            month.strftime("%Y-%m-%d") if row["needs_actual"] else "",
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=offset, column=index, value=value)
            if index == 6:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    for index, width in enumerate((14, 26, 22, 46, 34, 44, 16, 16, 12), start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = f"A{header + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, len(rows)


def _default_month():
    """The month a figure is for: the one that has CLOSED."""
    today = datetime.date.today()
    first = today.replace(day=1)
    return (first - datetime.timedelta(days=1)).replace(day=1)


def _number(value):
    if value in (None, "", "n/a", "N/A", "-"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("%", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def read_upload(stream):
    """Read the filled template → ``(entries, problems)``.

    Nothing is written. Every row is checked against the roster and the seeded
    cards, because a figure against a sales code nobody holds, or against a KPI
    that is not on that person's card, is silently invisible afterwards - the
    worst kind of import.
    """
    from openpyxl import load_workbook

    from .scorecard_automation.models import ScRoleKpiMapping

    book = load_workbook(stream, data_only=True, read_only=True)
    sheet = book[book.sheetnames[0]]

    header_row, headers = None, {}
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        labels = [str(c).strip().lower() if c is not None else "" for c in row]
        if "sales_code" in labels and "kpi_code" in labels:
            header_row = index
            headers = {name: position for position, name in enumerate(labels)}
            break
    if header_row is None:
        return [], ["That file has no sales_code / kpi_code header row. "
                    "Download the template and fill that in."]

    # Who holds which card, so a figure cannot be filed against the wrong
    # line. Built from the same three rosters _roster() uses, because a figure
    # for a teller would otherwise be refused for not being on the SALES
    # roster - which a teller never is.
    roles = {code: role for code, _name, _title, role in _roster()}
    on_card = {}
    for mapping in ScRoleKpiMapping.objects.all():
        on_card.setdefault(mapping.role_code, set()).add(mapping.kpi_code)

    entries, problems = [], []
    for index, row in enumerate(sheet.iter_rows(min_row=header_row + 1,
                                                values_only=True),
                                start=header_row + 1):
        def column(name):
            position = headers.get(name)
            return row[position] if position is not None and position < len(row) else None

        code = str(column("sales_code") or "").strip().upper()
        kpi = str(column("kpi_code") or "").strip()
        if not code and not kpi:
            continue
        if not code or not kpi:
            problems.append(f"Row {index}: needs both a sales code and a KPI code.")
            continue
        if code not in roles:
            problems.append(
                f"Row {index}: {code} is not on the DMC roster with a role that "
                f"has a card, so a figure against it would never be read.")
            continue
        if kpi not in on_card.get(roles[code], set()):
            problems.append(
                f"Row {index}: {kpi} is not on the {roles[code]} card, which is "
                f"the card {code} is measured on.")
            continue

        target = _number(column("target"))
        actual = _number(column("actual"))
        if target is None and actual is None:
            continue
        if target is not None and kpi not in MANUAL_TARGETS:
            problems.append(
                f"Row {index}: {kpi} takes its target from the DMC roster or "
                f"the role's card, so a target typed here would be ignored. "
                f"Leave it blank.")
            target = None
        if actual is not None and kpi not in MANUAL_ACTUALS:
            problems.append(
                f"Row {index}: {kpi} is computed from the warehouse, so an "
                f"actual typed here would never be used. Leave it blank.")
            actual = None
        if target is None and actual is None:
            continue

        month = _default_month()
        raw_month = column("month")
        if raw_month:
            parsed = _month_of(raw_month)
            if parsed is None:
                problems.append(
                    f"Row {index}: {raw_month!r} is not a date. Use YYYY-MM-DD.")
                continue
            month = parsed

        entries.append({"sales_code": code, "kpi_code": kpi, "target": target,
                        "actual": actual, "month": month, "row": index})
    return entries, problems


def _month_of(value):
    if isinstance(value, datetime.datetime):
        return value.date().replace(day=1)
    if isinstance(value, datetime.date):
        return value.replace(day=1)
    text = str(value).strip()[:10]
    for shape in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(text, shape).date().replace(day=1)
        except ValueError:
            continue
    return None


def apply_upload(entries, who=""):
    """Write the read entries → a count of what changed.

    Upserts: one figure per person per KPI (per month, for an actual), so
    uploading a correction replaces the figure rather than stacking a second
    one beside it. Who changed what is in the simple_history tables.
    """
    from .scorecard_automation.models import (
        ScEmployeeKpiTarget, ScEmployeePerformanceActual)

    targets = actuals = 0
    for entry in entries:
        if entry["target"] is not None:
            ScEmployeeKpiTarget.objects.update_or_create(
                sales_code=entry["sales_code"], kpi_code=entry["kpi_code"],
                year=entry["month"].year,
                defaults={"kpi_target": entry["target"],
                          "source": "Administration upload",
                          "updated_by": who[:150]})
            targets += 1
        if entry["actual"] is not None:
            # ScEmployeePerformanceActual.save refuses an UPDATE without a
            # change reason, so one is set before the row is written rather
            # than afterwards - a correction to somebody's figure is exactly
            # the change that model wants a reason for.
            row, created = ScEmployeePerformanceActual.objects.get_or_create(
                sales_code=entry["sales_code"], kpi_code=entry["kpi_code"],
                eom_date=entry["month"],
                defaults={"kpi_value": entry["actual"]})
            if not created:
                row.kpi_value = entry["actual"]
                reason = f"Administration upload by {who or 'an administrator'}"
                row._change_reason = reason[:100]
                row.save()
            actuals += 1
    return {"targets": targets, "actuals": actuals}
