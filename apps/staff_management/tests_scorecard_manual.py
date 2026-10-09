"""Administration loading the figures the system cannot compute.

NPS comes from the customer survey, training hours from the learning system,
leave from HR, the audit score from Internal Audit. None of those is in this
platform and none of them is going to be, so they get typed in - and the rules
around typing them in are the whole point of these tests:

* only an administrator may;
* a figure against a sales code nobody holds, or against a KPI that is not on
  that person's card, is refused rather than written somewhere nothing reads;
* nothing is saved until a second call says to save it;
* a typed figure can never overwrite a measured one.
"""

import datetime
from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.portfolio.models import Profile
from .live_scorecard import build_card
from .models import BranchEmployeeDmcData
from .scorecard_manual import (
    MANUAL_ACTUALS, MANUAL_TARGETS, _default_month, build_template,
    rows_needing_a_figure,
)

BASE = "/staff_management/scorecard-automation/manual-figures/"
XLSX = ("application/vnd.openxmlformats-officedocument"
        ".spreadsheetml.sheet")


def workbook(rows):
    """A filled template, as Administration would send it back."""
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="a note the reader has to skip")
    headers = ["sales_code", "staff_name", "staff_role", "kpi_code",
               "kpi_name", "what_this_is", "target", "actual", "month"]
    for index, name in enumerate(headers, start=1):
        ws.cell(row=4, column=index, value=name)
    for offset, row in enumerate(rows, start=5):
        for index, name in enumerate(headers, start=1):
            ws.cell(row=offset, column=index, value=row.get(name))
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile("figures.xlsx", buffer.read(), content_type=XLSX)


class ManualFigureTests(TestCase):
    CODE = "VO4300"

    def setUp(self):
        BranchEmployeeDmcData.objects.create(
            staff_pf_number=4300, staff_name="Vincent Ogare",
            sales_code=self.CODE, staff_role="PB RM", staff_branch="Rehani",
            active=1, target_new_customers=32)

        self.admin = User.objects.create_user("admin1", password="x",
                                              is_superuser=True,
                                              first_name="Ada", last_name="Min")
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

        self.rm = User.objects.create_user("vogare", password="x")
        Profile.objects.update_or_create(
            user=self.rm, defaults={"sales_code": self.CODE})
        self.rm_client = APIClient()
        self.rm_client.force_authenticate(self.rm)

    # ── who may ──────────────────────────────────────────────────────────
    def test_an_rm_cannot_see_everybodys_figures(self):
        for url in (BASE, BASE + "template/"):
            with self.subTest(url):
                self.assertEqual(self.rm_client.get(url).status_code, 403)

    def test_a_merely_is_staff_account_cannot_set_a_figure(self):
        """is_staff is legacy noise in this estate - a great many accounts
        carry it for reasons that have nothing to do with scorecards. A screen
        that sets what people are paid on needs a real signal."""
        staffish = User.objects.create_user("staffish", password="x",
                                            is_staff=True)
        client = APIClient()
        client.force_authenticate(staffish)
        self.assertEqual(client.get(BASE).status_code, 403)
        self.assertEqual(
            client.post(BASE + "entry/",
                        {"sales_code": self.CODE, "kpi_code": "nps",
                         "actual": 0.8}, format="json").status_code, 403)

    def test_an_administration_group_may(self):
        from django.contrib.auth.models import Group

        member = User.objects.create_user("admingroup", password="x")
        member.groups.add(Group.objects.create(name="staff_mgt"))
        client = APIClient()
        client.force_authenticate(member)
        self.assertEqual(client.get(BASE).status_code, 200)

    def test_an_rm_cannot_set_a_figure(self):
        response = self.rm_client.post(
            BASE + "entry/",
            {"sales_code": self.CODE, "kpi_code": "nps", "actual": 0.8},
            format="json")
        self.assertEqual(response.status_code, 403)

    # ── what needs filling in ────────────────────────────────────────────
    def test_the_outstanding_list_is_built_from_the_roster_and_the_cards(self):
        response = self.client.get(BASE)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertGreater(response.data["total"], 0)
        self.assertEqual(response.data["loaded"], 0)
        self.assertEqual(response.data["outstanding"], response.data["total"])

        codes = {row["kpi_code"] for row in response.data["rows"]}
        self.assertIn("nps", codes, "the PB card's NPS line needs a figure")
        for code in codes:
            with self.subTest(code):
                self.assertTrue(code in MANUAL_ACTUALS or code in MANUAL_TARGETS,
                                f"{code} is listed but the system can compute it")

    def test_only_lines_on_that_persons_card_are_listed(self):
        """A template listing KPIs nobody is measured on is a template people
        stop reading."""
        codes = {row["kpi_code"] for row in rows_needing_a_figure(self.CODE)}
        self.assertIn("nps", codes)
        self.assertNotIn("trade_income", codes,
                         "trade income is not on the PB card")

    def test_the_template_is_prefilled_with_who_and_which_line(self):
        buffer, count = build_template(self.CODE)
        self.assertGreater(count, 0)
        sheet = load_workbook(BytesIO(buffer.read())).active
        text = "\n".join(
            " ".join("" if c.value is None else str(c.value) for c in row)
            for row in sheet.iter_rows())
        self.assertIn(self.CODE, text)
        self.assertIn("Vincent Ogare", text)
        self.assertIn("nps", text)
        self.assertIn("sales_code", text)

    def test_the_template_downloads(self):
        response = self.client.get(BASE + "template/")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertTrue(response.content.startswith(b"PK"))

    # ── reading before writing ───────────────────────────────────────────
    def test_a_first_upload_reports_and_writes_nothing(self):
        from .scorecard_automation.models import ScEmployeePerformanceActual

        response = self.client.post(
            BASE + "upload/",
            {"file": workbook([{"sales_code": self.CODE, "kpi_code": "nps",
                                "actual": 0.72}])},
            format="multipart")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertFalse(response.data["applied"])
        self.assertEqual(response.data["found"], 1)
        self.assertEqual(ScEmployeePerformanceActual.objects.count(), 0,
                         "a first upload must write nothing")

    def test_a_second_call_with_apply_saves(self):
        from .scorecard_automation.models import ScEmployeePerformanceActual

        response = self.client.post(
            BASE + "upload/",
            {"file": workbook([{"sales_code": self.CODE, "kpi_code": "nps",
                                "actual": 0.72}]),
             "apply": "true"},
            format="multipart")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.data["applied"])
        self.assertEqual(response.data["written"]["actuals"], 1)

        row = ScEmployeePerformanceActual.objects.get(sales_code=self.CODE)
        self.assertEqual(row.kpi_code, "nps")
        self.assertAlmostEqual(row.kpi_value, 0.72)
        self.assertEqual(row.eom_date, _default_month())

    def test_a_correction_replaces_rather_than_stacking(self):
        from .scorecard_automation.models import ScEmployeePerformanceActual

        for value in (0.6, 0.81):
            self.client.post(
                BASE + "upload/",
                {"file": workbook([{"sales_code": self.CODE,
                                    "kpi_code": "nps", "actual": value}]),
                 "apply": "true"},
                format="multipart")
        rows = ScEmployeePerformanceActual.objects.filter(sales_code=self.CODE)
        self.assertEqual(rows.count(), 1)
        self.assertAlmostEqual(rows.first().kpi_value, 0.81)

    # ── what is refused ──────────────────────────────────────────────────
    def test_a_figure_against_an_unknown_sales_code_is_refused(self):
        response = self.client.post(
            BASE + "upload/",
            {"file": workbook([{"sales_code": "ZZ9999", "kpi_code": "nps",
                                "actual": 0.5}]), "apply": "true"},
            format="multipart")
        self.assertEqual(response.data["found"], 0)
        self.assertTrue(any("not on the DMC roster" in p
                            for p in response.data["problems"]))

    def test_a_figure_against_a_kpi_not_on_that_card_is_refused(self):
        response = self.client.post(
            BASE + "upload/",
            {"file": workbook([{"sales_code": self.CODE,
                                "kpi_code": "trade_income", "actual": 10}]),
             "apply": "true"},
            format="multipart")
        self.assertEqual(response.data["found"], 0)
        self.assertTrue(any("is not on the pb_rm card" in p
                            for p in response.data["problems"]))

    def test_an_actual_for_a_line_the_warehouse_computes_is_refused(self):
        """Accepting it would be worse than refusing it: the card would never
        read the figure, and whoever typed it would believe it had."""
        response = self.client.post(
            BASE + "upload/",
            {"file": workbook([{"sales_code": self.CODE,
                                "kpi_code": "drawdowns", "actual": 5_000_000}]),
             "apply": "true"},
            format="multipart")
        self.assertEqual(response.data["found"], 0)
        self.assertTrue(any("computed from the warehouse" in p
                            for p in response.data["problems"]))

    def test_a_file_with_no_header_row_says_so(self):
        wb = Workbook()
        wb.active.cell(row=1, column=1, value="nothing useful")
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = self.client.post(
            BASE + "upload/",
            {"file": SimpleUploadedFile("x.xlsx", buffer.read(),
                                        content_type=XLSX)},
            format="multipart")
        self.assertEqual(response.data["found"], 0)
        self.assertTrue(any("no sales_code" in p
                            for p in response.data["problems"]))

    # ── one figure at a time ─────────────────────────────────────────────
    def test_a_single_entry_saves(self):
        response = self.client.post(
            BASE + "entry/",
            {"sales_code": self.CODE, "kpi_code": "nps", "actual": 0.65},
            format="json")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.data["written"]["actuals"], 1)

    def test_a_single_entry_for_the_wrong_card_is_refused(self):
        response = self.client.post(
            BASE + "entry/",
            {"sales_code": self.CODE, "kpi_code": "trade_income", "actual": 1},
            format="json")
        self.assertEqual(response.status_code, 400)


class LoadedFiguresReachTheCardTests(TestCase):
    """The point of all of it: a figure typed in shows up scored."""

    CODE = "VO4301"

    def setUp(self):
        BranchEmployeeDmcData.objects.create(
            staff_pf_number=4301, staff_name="Vincent Ogare",
            sales_code=self.CODE, staff_role="PB RM", staff_branch="Rehani",
            active=1)

    def load(self, kpi, value, month=None):
        from .scorecard_automation.models import ScEmployeePerformanceActual

        ScEmployeePerformanceActual.objects.create(
            sales_code=self.CODE, kpi_code=kpi,
            eom_date=month or _default_month(), kpi_value=value)

    def line(self, code):
        card = build_card(self.CODE)
        self.assertTrue(card["has_card"], card)
        return {ln["kpi_code"]: ln
                for g in card["perspectives"] for ln in g["lines"]}[code]

    def test_an_unloaded_line_says_it_is_measured_elsewhere(self):
        nps = self.line("nps")
        self.assertIsNone(nps["ytd_actual"])
        self.assertEqual(nps["pending_label"], "Not measured here")

    def test_a_loaded_figure_is_scored_against_the_role_threshold(self):
        """NPS: 60% is the target on every card that carries it."""
        self.load("nps", 0.72)
        nps = self.line("nps")
        self.assertAlmostEqual(nps["ytd_actual"], 0.72)
        self.assertEqual(nps["annual_target"], 0.6)
        self.assertAlmostEqual(nps["score"], 1.2)
        self.assertFalse(nps["pending"])

    def test_the_card_says_the_figure_was_typed_in(self):
        """A hand-typed figure must never be mistaken for a measured one."""
        self.load("nps", 0.5)
        self.assertIn("loaded by Administration", self.line("nps")["as_at"])

    def test_a_figure_for_an_older_month_is_not_carried_forward(self):
        """A survey score from four months ago presented as this month's is
        worse than a blank."""
        old = _default_month() - datetime.timedelta(days=120)
        self.load("nps", 0.9, month=old.replace(day=1))
        nps = self.line("nps")
        self.assertIsNone(nps["ytd_actual"])

    def test_a_typed_figure_cannot_overwrite_a_measured_one(self):
        """The warehouse always wins, so a stale upload cannot quietly replace
        a live number."""
        self.load("drawdowns", 999_999_999)
        drawdowns = self.line("drawdowns")
        self.assertNotEqual(drawdowns["ytd_actual"], 999_999_999)
