"""Operating costs, against the Finance team's own rules.

Every rule here comes out of ``cost_analysis_sap_report.ipynb``, and each one
changes the answer, so each gets a test:

* a credit is NEGATIVE — without it every total is out by twice the credits;
* the date is not on the posting, it comes off the assignment number with the
  reference document number as a fallback;
* an operating expense is a 7-series GL minus three specific accounts;
* an unmapped cost centre is DROPPED, not bucketed;
* a branch is grouped by NAME, because several branches have two cost centres.

Plus the thing that makes the whole feature safe: a load replaces its snapshot,
because each SAP extract is a full copy rather than an increment.
"""

import datetime
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from . import cost_analysis, cost_ingest
from .models import (
    OperatingExpenseMapping, SapCostCentre, SapCostExclusion, SapCostPosting,
    SapSubsidiary,
)

BASE = "/staff_management/costs/"
XLSX = ("application/vnd.openxmlformats-officedocument"
        ".spreadsheetml.sheet")
SNAPSHOT = datetime.date(2026, 9, 10)
THIS_YEAR = datetime.date.today().year


def posting(**kwargs):
    """A posting with the fields the report needs, defaults that pass."""
    defaults = dict(
        company_code=1000, document_number="DOC1", item_number="1",
        fiscal_year=THIS_YEAR, debit_credit_indicator="S",
        cost_centre=110001, gl_account_number="705070001",
        lcy_amount=1000, amount=1000,
        entry_date=datetime.date(THIS_YEAR, 3, 15), entry_month=f"{THIS_YEAR}-03",
        snapshot_date=SNAPSHOT)
    defaults.update(kwargs)
    return SapCostPosting.objects.create(**defaults)


class TransformRuleTests(TestCase):
    """The notebook's arithmetic, before any database is involved."""

    def test_a_credit_is_negative(self):
        """SAP gives the magnitude either way and SHKZG carries the sign.
        Without this, a credit ADDS to a cost instead of reducing it."""
        self.assertEqual(cost_ingest.amount_for("1,234.00", "S"), 1234.0)
        self.assertEqual(cost_ingest.amount_for("1,234.00", "H"), -1234.0)
        self.assertEqual(cost_ingest.amount_for("1,234.00", "h"), -1234.0,
                         "the indicator is not reliably upper case")

    def test_a_trailing_minus_is_read(self):
        """SAP writes some credits as 1,234.00- rather than -1,234.00."""
        self.assertEqual(cost_ingest.amount_for("1,234.00-", "S"), -1234.0)

    def test_no_amount_is_none_rather_than_nought(self):
        """A posting with no amount is not a posting of nothing."""
        for value in (None, "", "   ", "n/a"):
            with self.subTest(repr(value)):
                self.assertIsNone(cost_ingest.amount_for(value, "S"))

    def test_the_date_comes_off_the_assignment_number(self):
        self.assertEqual(
            cost_ingest.entry_date_for("20260315", "00000000"),
            datetime.date(2026, 3, 15))

    def test_it_falls_back_to_the_reference_document(self):
        """Neither field is reliably populated on its own."""
        self.assertEqual(
            cost_ingest.entry_date_for("00000000", "20260215"),
            datetime.date(2026, 2, 15))

    def test_an_impossible_date_falls_through_rather_than_raising(self):
        """The notebook's errors='coerce'. 31 September is not a date."""
        self.assertEqual(
            cost_ingest.entry_date_for("20260931", "20260215"),
            datetime.date(2026, 2, 15))

    def test_no_usable_date_at_all_is_none(self):
        self.assertIsNone(cost_ingest.entry_date_for("", "0"))
        self.assertIsNone(cost_ingest.entry_date_for("00000000", "00000000"))

    def test_the_gl_first_digit_classification(self):
        for gl, expected in (("170150001", "assets"),
                             ("230000023", "liabilities"),
                             ("705070001", "other expenses"),
                             ("405000001", "interest income"),
                             ("", "unmapped")):
            with self.subTest(gl):
                self.assertEqual(cost_analysis.gl_type(gl), expected)

    def test_the_snapshot_is_read_off_the_file_name(self):
        """The extract carries its date in its name, as dd-mm-yyyy or
        dd-mm-yy - the same two shapes the Trino query tries."""
        self.assertEqual(
            cost_ingest._snapshot_from(None, "cost_data_10-09-2026.xls"),
            datetime.date(2026, 9, 10))
        self.assertEqual(
            cost_ingest._snapshot_from(None, "cost_data_10-09-26.xls"),
            datetime.date(2026, 9, 10))

    def test_the_query_still_takes_only_the_newest_extract(self):
        """If that WHERE is ever dropped, every posting is counted once per
        extract that ever landed."""
        self.assertIn("WHERE d.dt = m.max_dt", cost_ingest.TRINO_QUERY)
        self.assertIn("delta.gold_db.sap_cost_data", cost_ingest.TRINO_QUERY)


class ReportFilterTests(TestCase):
    """What counts as an operating expense, and what is dropped."""

    def setUp(self):
        # All three are already seeded by migration 0034; this asserts the
        # shape the tests below rely on rather than inventing it.
        for centre, name, dept, second, reportable in (
            (110001, "ICT", "ICT", "", True),
            (510, "NYERI", "RETAIL", "BRANCH BUSINESS", True),
            (40001, "BRANCH GROWTH", "RETAIL", "", False),
        ):
            SapCostCentre.objects.update_or_create(
                cost_centre=centre,
                defaults={"cost_centre_name": name, "cost_department": dept,
                          "cost_department_2": second,
                          "is_branch": second == "BRANCH BUSINESS",
                          "is_reportable": reportable})
        OperatingExpenseMapping.objects.update_or_create(
            gl="705070001",
            defaults={"expense_type": "Legal and professional fees",
                      "operating_expense": "Legal fees",
                      "actual_gl_name": "Legal and professional fees"})

    def amounts(self):
        return {row["department"]: float(row["amount"] or 0)
                for row in cost_analysis.by_department(SNAPSHOT)}

    def test_a_seven_series_posting_with_a_mapping_counts(self):
        posting(amount=5000)
        self.assertEqual(self.amounts(), {"ICT": 5000.0})

    def test_a_non_seven_series_posting_does_not(self):
        """Only the 7-series is an operating expense. A liability posting on
        230000023 is not a cost."""
        OperatingExpenseMapping.objects.update_or_create(
            gl="230000023",
            defaults={"expense_type": "Staff costs",
                      "operating_expense": "Staff costs",
                      "actual_gl_name": "Staff costs"})
        posting(gl_account_number="230000023", amount=9999)
        self.assertEqual(self.amounts(), {})

    def test_an_excluded_gl_does_not(self):
        """Loan loss charges, lease liability interest and interest
        adjustments sit on 7-series accounts and are not operating expenses.
        Counting them would inflate every department."""
        SapCostExclusion.objects.update_or_create(
            gl="705190002", defaults={"reason": "loan loss"})
        OperatingExpenseMapping.objects.update_or_create(
            gl="705190002",
            defaults={"expense_type": "Operating Expenses",
                      "operating_expense": "Loan loss",
                      "actual_gl_name": "Loan loss charges"})
        posting(gl_account_number="705190002", amount=7777)
        self.assertEqual(self.amounts(), {})

    def test_a_gl_with_no_mapping_row_does_not(self):
        posting(gl_account_number="709999999", amount=4321)
        self.assertEqual(self.amounts(), {})

    def test_an_unmapped_cost_centre_is_dropped_not_bucketed(self):
        """A cost with nowhere to go must not land in a department that did
        not spend it."""
        posting(cost_centre=999999, amount=1234)
        self.assertEqual(self.amounts(), {})

    def test_branch_growth_is_dropped(self):
        """A holding cost centre rather than a place."""
        posting(cost_centre=40001, amount=2222)
        self.assertEqual(self.amounts(), {})

    def test_a_posting_with_no_month_is_dropped_from_the_monthly_view(self):
        """It is still in the health view - it is money that exists - but it
        cannot be put in a month."""
        posting(entry_date=None, entry_month="", amount=3333)
        self.assertEqual(self.amounts(), {})
        self.assertEqual(cost_analysis.health(SNAPSHOT)["postings_without_a_date"], 1)

    def test_credits_reduce_the_total(self):
        posting(amount=10_000)
        posting(document_number="DOC2", debit_credit_indicator="H",
                lcy_amount=4_000, amount=-4_000)
        self.assertEqual(self.amounts(), {"ICT": 6000.0})

    def test_only_the_named_snapshot_is_read(self):
        """Each extract is a full copy, so reading across snapshots would
        count the same posting twice."""
        posting(amount=1000)
        posting(document_number="OLD", snapshot_date=datetime.date(2026, 8, 10),
                amount=9999)
        self.assertEqual(self.amounts(), {"ICT": 1000.0})


class BranchAndLineTests(TestCase):
    def setUp(self):
        # Kisumu has two cost centres - 18 and 600 - and so do Eldoret,
        # Hurlingham and Mombasa. Grouping on the CODE would report each twice
        # at half its cost.
        for centre in (18, 600):
            SapCostCentre.objects.update_or_create(
                cost_centre=centre,
                defaults={"cost_centre_name": "KISUMU",
                          "cost_department": "RETAIL",
                          "cost_department_2": "BRANCH BUSINESS",
                          "is_branch": True})
        SapCostCentre.objects.update_or_create(
            cost_centre=110001,
            defaults={"cost_centre_name": "ICT", "cost_department": "ICT",
                      "cost_department_2": "", "is_branch": False})
        OperatingExpenseMapping.objects.update_or_create(
            gl="705070001",
            defaults={"expense_type": "Legal and professional fees",
                      "operating_expense": "Legal fees",
                      "actual_gl_name": "Legal and professional fees"})
        OperatingExpenseMapping.objects.update_or_create(
            gl="705150001",
            defaults={"expense_type": "Operating Expenses",
                      "operating_expense": "DPF",
                      "actual_gl_name": "Deposit protection fund"})
        SapSubsidiary.objects.update_or_create(
            company_code=1000, defaults={"subsidiary": "HFCB"})

    def test_a_branch_with_two_cost_centres_is_reported_once(self):
        posting(cost_centre=18, amount=3000)
        posting(cost_centre=600, document_number="D2", amount=4000)
        rows = cost_analysis.by_branch(SNAPSHOT)
        self.assertEqual(len(rows), 1, rows)
        self.assertEqual(rows[0]["branch"], "KISUMU")
        self.assertEqual(float(rows[0]["amount"]), 7000.0)

    def test_a_head_office_cost_centre_is_not_a_branch(self):
        posting(cost_centre=110001, amount=5000)
        self.assertEqual(cost_analysis.by_branch(SNAPSHOT), [])

    def test_the_expense_line_is_the_actual_gl_name(self):
        """That is what gives the ten names the CEO deck had hard-coded -
        "Deposit protection fund", "Legal and professional fees" - all of which
        were already in the GL mapping."""
        posting(amount=1000)
        posting(gl_account_number="705150001", document_number="D2",
                amount=2000)
        lines = {row["expense_line"]: float(row["amount"])
                 for row in cost_analysis.by_expense_line(SNAPSHOT)}
        self.assertEqual(lines, {"Legal and professional fees": 1000.0,
                                 "Deposit protection fund": 2000.0})

    def test_the_subsidiary_split_names_an_unmapped_company_code(self):
        """A department total that silently mixed the subsidiaries into the
        bank would overstate the bank."""
        posting(amount=1000)
        # 2000 is seeded as HFCB-PROPERTIES; 9000 is in no mapping, which is
        # the case that must not silently land in the bank's total.
        posting(company_code=2000, document_number="D2", amount=500)
        posting(company_code=9000, document_number="D3", amount=250)
        rows = {row["subsidiary"]: float(row["amount"])
                for row in cost_analysis.by_subsidiary(SNAPSHOT)}
        self.assertEqual(rows["HFCB"], 1000.0)
        self.assertEqual(rows["HFCB-PROPERTIES"], 500.0)
        self.assertEqual(rows["company 9000"], 250.0,
                         "an unmapped company code must be named, not merged")


class HealthTests(TestCase):
    """What the load cannot place is the one failure that silently shrinks
    every department total."""

    def setUp(self):
        SapCostCentre.objects.update_or_create(
            cost_centre=110001,
            defaults={"cost_centre_name": "ICT", "cost_department": "ICT",
                      "cost_department_2": "", "is_branch": False})
        OperatingExpenseMapping.objects.update_or_create(
            gl="705070001",
            defaults={"expense_type": "Legal", "operating_expense": "Legal",
                      "actual_gl_name": "Legal and professional fees"})

    def test_nothing_loaded_says_so_rather_than_returning_zeroes(self):
        health = cost_analysis.health()
        self.assertIsNone(health["snapshot_date"])
        self.assertEqual(health["postings"], 0)

    def test_an_unmapped_cost_centre_is_named_with_what_it_carries(self):
        posting(cost_centre=777777, amount=50_000)
        health = cost_analysis.health(SNAPSHOT)
        missing = {row["cost_centre"]: row["amount"]
                   for row in health["unmapped_cost_centres"]}
        self.assertEqual(missing, {777777: 50_000.0})
        self.assertEqual(health["reportable_amount"], 0.0)

    def test_an_unmapped_gl_is_named_too(self):
        posting(gl_account_number="707777777", amount=1_500)
        health = cost_analysis.health(SNAPSHOT)
        self.assertEqual([row["gl"] for row in health["unmapped_gls"]],
                         ["707777777"])


class LoadTests(TestCase):
    def rows(self, **overrides):
        row = {
            "company_code": 1000, "document_number": "DOC1", "item_number": "1",
            "posting_key": "40", "cost_element_category": "S",
            "debit_credit_indicator": "S", "lcy_amount": 1000,
            "fiscal_year": THIS_YEAR, "assignment_number": "20260315",
            "reference_document_number": "00000000",
            "cost_center": 110001, "gl_account_number": "705070001",
            "source_file": "cost_data_10-09-2026", "dt": SNAPSHOT,
        }
        row.update(overrides)
        return [row]

    def test_a_load_computes_the_month_and_the_signed_amount(self):
        postings, report = cost_ingest.rows_to_postings(self.rows())
        self.assertEqual(report["kept"], 1)
        self.assertEqual(float(postings[0].amount), 1000.0)
        self.assertEqual(postings[0].entry_month, f"{THIS_YEAR}-03")
        self.assertEqual(postings[0].snapshot_date, SNAPSHOT)

    def test_a_row_with_no_gl_is_skipped_and_counted(self):
        _, report = cost_ingest.rows_to_postings(
            self.rows(gl_account_number=""))
        self.assertEqual(report["kept"], 0)
        self.assertEqual(report["skipped_no_gl"], 1)

    def test_a_load_replaces_its_snapshot_rather_than_adding_to_it(self):
        """Each SAP extract is a full copy. Appending would double every
        posting the next time the same extract arrived."""
        for _ in range(2):
            postings, report = cost_ingest.rows_to_postings(self.rows())
            result = cost_ingest.save_postings(postings, report)
        self.assertEqual(SapCostPosting.objects.count(), 1)
        self.assertEqual(result["replaced"], 1)

    def test_two_extract_dates_in_one_file_are_refused(self):
        """Replacing "the" snapshot would then delete half of what was just
        read."""
        rows = self.rows() + self.rows(dt=datetime.date(2026, 8, 10),
                                       document_number="DOC2",
                                       source_file="cost_data_10-08-2026")
        postings, report = cost_ingest.rows_to_postings(rows)
        result = cost_ingest.save_postings(postings, report)
        self.assertEqual(result["written"], 0)
        self.assertIn("one at a time", result["detail"])
        self.assertEqual(SapCostPosting.objects.count(), 0)


class EndpointTests(TestCase):
    def setUp(self):
        SapCostCentre.objects.update_or_create(
            cost_centre=110001,
            defaults={"cost_centre_name": "ICT", "cost_department": "ICT",
                      "cost_department_2": "", "is_branch": False})
        OperatingExpenseMapping.objects.update_or_create(
            gl="705070001",
            defaults={"expense_type": "Legal", "operating_expense": "Legal",
                      "actual_gl_name": "Legal and professional fees"})
        self.admin = User.objects.create_user("costadmin", password="x",
                                              is_superuser=True)
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

        self.reader = User.objects.create_user("costreader", password="x")
        self.read_client = APIClient()
        self.read_client.force_authenticate(self.reader)

    def workbook(self, rows):
        book = Workbook()
        sheet = book.active
        sheet.append(["An extract from SAP"])
        headers = list(rows[0].keys())
        sheet.append(headers)
        for row in rows:
            sheet.append([row[name] for name in headers])
        buffer = BytesIO()
        book.save(buffer)
        buffer.seek(0)
        return SimpleUploadedFile("extract.xlsx", buffer.read(),
                                  content_type=XLSX)

    def test_the_summary_says_when_nothing_is_loaded(self):
        response = self.client.get(BASE + "summary/")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIsNone(response.data["snapshot_date"])
        self.assertIn("No SAP cost extract has been loaded",
                      response.data["detail"])

    def test_anybody_signed_in_can_read_the_summary(self):
        """The CEO deck reads it."""
        posting(amount=1000)
        response = self.read_client.get(BASE + "summary/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            [row["department"] for row in response.data["by_department"]],
            ["ICT"])

    def test_a_reader_cannot_load_an_extract(self):
        response = self.read_client.post(
            BASE + "upload/", {"file": self.workbook([{
                "lcy_amount": 1, "gl_account_number": "705070001",
                "cost_center": 110001, "fiscal_year": THIS_YEAR,
                "assignment_number": "20260315",
                "debit_credit_indicator": "S", "dt": "2026-09-10"}])},
            format="multipart")
        self.assertEqual(response.status_code, 403)

    def test_a_reader_cannot_change_a_cost_centre(self):
        centre = SapCostCentre.objects.get(cost_centre=110001)
        response = self.read_client.patch(
            f"{BASE}cost-centres/{centre.id}/",
            {"cost_department": "FINANCE"}, format="json")
        self.assertEqual(response.status_code, 403)

    def test_an_administrator_can_change_a_cost_centre(self):
        centre = SapCostCentre.objects.get(cost_centre=110001)
        response = self.client.patch(
            f"{BASE}cost-centres/{centre.id}/",
            {"cost_department": "FINANCE"}, format="json")
        self.assertEqual(response.status_code, 200, response.content)
        centre.refresh_from_db()
        self.assertEqual(centre.cost_department, "FINANCE")

    def test_marking_a_cost_centre_a_branch_is_derived_not_typed(self):
        """One answer to "is this a branch": the second department level the
        Finance mapping uses."""
        centre = SapCostCentre.objects.get(cost_centre=110001)
        self.client.patch(f"{BASE}cost-centres/{centre.id}/",
                          {"cost_department_2": "BRANCH BUSINESS"},
                          format="json")
        centre.refresh_from_db()
        self.assertTrue(centre.is_branch)

    def test_an_upload_reads_and_reports_before_it_writes(self):
        rows = [{"lcy_amount": 2500, "gl_account_number": "705070001",
                 "cost_center": 110001, "fiscal_year": THIS_YEAR,
                 "assignment_number": "20260315",
                 "debit_credit_indicator": "S", "dt": "2026-09-10"}]
        response = self.client.post(BASE + "upload/",
                                    {"file": self.workbook(rows)},
                                    format="multipart")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertFalse(response.data["applied"])
        self.assertEqual(response.data["report"]["kept"], 1)
        self.assertEqual(SapCostPosting.objects.count(), 0)

    def test_a_second_call_with_apply_writes_it(self):
        rows = [{"lcy_amount": 2500, "gl_account_number": "705070001",
                 "cost_center": 110001, "fiscal_year": THIS_YEAR,
                 "assignment_number": "20260315",
                 "debit_credit_indicator": "S", "dt": "2026-09-10"}]
        response = self.client.post(
            BASE + "upload/",
            {"file": self.workbook(rows), "apply": "true"},
            format="multipart")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.data["applied"])
        self.assertEqual(SapCostPosting.objects.count(), 1)
        self.assertEqual(float(SapCostPosting.objects.get().amount), 2500.0)

    def test_an_empty_upload_does_not_wipe_what_is_there(self):
        posting(amount=1000)
        book = Workbook()
        book.active.append(["lcy_amount", "gl_account_number"])
        buffer = BytesIO()
        book.save(buffer)
        buffer.seek(0)
        response = self.client.post(
            BASE + "upload/",
            {"file": SimpleUploadedFile("empty.xlsx", buffer.read(),
                                        content_type=XLSX),
             "apply": "true"},
            format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(SapCostPosting.objects.count(), 1)

    def test_a_file_with_no_recognisable_header_says_so(self):
        book = Workbook()
        book.active.append(["nothing", "useful"])
        buffer = BytesIO()
        book.save(buffer)
        buffer.seek(0)
        response = self.client.post(
            BASE + "upload/",
            {"file": SimpleUploadedFile("x.xlsx", buffer.read(),
                                        content_type=XLSX)},
            format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertIn("header row", str(response.data))

    def test_the_snapshots_list_shows_what_has_been_loaded(self):
        posting(amount=1000)
        response = self.client.get(BASE + "snapshots/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["snapshots"][0]["snapshot_date"],
                         SNAPSHOT.isoformat())
        self.assertEqual(response.data["snapshots"][0]["postings"], 1)

    def test_the_postings_behind_a_figure_can_be_listed(self):
        posting(amount=1000)
        response = self.client.get(
            BASE + f"postings/?cost_centre=110001&entry_month={THIS_YEAR}-03")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)

    def test_a_finance_group_member_may_change_a_mapping(self):
        member = User.objects.create_user("fin", password="x")
        member.groups.add(Group.objects.create(name="finance"))
        client = APIClient()
        client.force_authenticate(member)
        centre = SapCostCentre.objects.get(cost_centre=110001)
        response = client.patch(f"{BASE}cost-centres/{centre.id}/",
                                {"cost_department": "ICT & Digital"},
                                format="json")
        self.assertEqual(response.status_code, 200, response.content)


class SeededMappingTests(TestCase):
    """The mapping the whole thing depends on, as migration 0034 seeds it."""

    def test_the_cost_centres_are_seeded(self):
        self.assertGreaterEqual(SapCostCentre.objects.count(), 80)

    def test_the_branches_are_marked(self):
        branches = set(SapCostCentre.objects.filter(is_branch=True)
                       .values_list("cost_centre_name", flat=True))
        for name in ("NYERI", "KISUMU", "MOMBASA", "REHANI", "ELDORET"):
            with self.subTest(name):
                self.assertIn(name, branches)

    def test_a_branch_with_two_codes_keeps_one_name(self):
        for name in ("KISUMU", "ELDORET", "HURLINGHAM", "MOMBASA"):
            with self.subTest(name):
                self.assertGreaterEqual(
                    SapCostCentre.objects.filter(
                        cost_centre_name=name).count(), 2,
                    f"{name} should have more than one cost centre")

    def test_branch_growth_is_not_reportable(self):
        """A holding cost centre rather than a place."""
        centre = SapCostCentre.objects.get(cost_centre_name="BRANCH GROWTH")
        self.assertFalse(centre.is_reportable)

    def test_the_three_exclusions_are_seeded_with_their_reasons(self):
        """In the notebook these are three numbers in an if, with a comment.
        Nobody outside it could see or change them."""
        exclusions = dict(SapCostExclusion.objects.values_list("gl", "reason"))
        self.assertEqual(set(exclusions), {"705190002", "701010014",
                                           "705190006"})
        for gl, reason in exclusions.items():
            with self.subTest(gl):
                self.assertTrue(reason.strip())

    def test_the_six_subsidiaries_are_seeded(self):
        subs = dict(SapSubsidiary.objects.values_list("company_code",
                                                      "subsidiary"))
        self.assertEqual(subs[1000], "HFCB")
        self.assertEqual(len(subs), 6)

    def test_the_ten_hardcoded_expense_lines_all_exist_in_the_gl_mapping(self):
        """The ten names the CEO deck had typed in as constants. If one of
        these is not in the mapping, that line silently disappears from the
        real report and somebody will say the new figures are wrong."""
        wanted = [
            "Staff costs", "IT Maintenance", "Marketing & Product Development",
            "Operational Losses", "Deposit protection fund",
            "Legal and professional fees", "Debt collection fees",
            "Lease Depreciation", "Depreciation Computer software",
            "Office Expenses",
        ]
        for name in wanted:
            with self.subTest(name):
                self.assertTrue(
                    OperatingExpenseMapping.objects.filter(
                        actual_gl_name__iexact=name).exists()
                    or OperatingExpenseMapping.objects.filter(
                        expense_type__iexact=name).exists(),
                    f"{name} is on the CEO deck but in no GL mapping row")
