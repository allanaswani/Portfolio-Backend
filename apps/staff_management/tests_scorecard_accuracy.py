"""The two things that have to be right before a card can be signed.

1. the TARGET is the person's own, off the per-person DMC table
2. the target is sliced to the same date as the ACTUAL it is compared with

Both have been wrong, and both were wrong in the direction that flatters or
punishes somebody by an order of magnitude, so each has its own test.
"""

import datetime

from django.contrib.auth.models import User
from django.test import TestCase
from rest_framework.test import APIClient

from apps.portfolio.models import Profile
from .live_scorecard import (
    STAFF_TABLE, BRANCH_TABLE, build_card, last_closed_month_end, prorate_to,
    roster_target_row, score_for,
)
from .models import BranchEmployeeDmcData, BranchFinalEmployeeDmcData
from .scorecard_signoff import build_card_workbook, period_for


class TargetsComeFromThePersonsOwnRowTests(TestCase):
    """``branch_employee_dmc_data`` is one row per SALES PERSON.
    ``branch_final_employee_dmc_data`` is one row per BRANCH. Reading the
    branch one for an RM scored an SME RM against a 14 million disbursement
    target where their card says 270 million."""

    CODE = "CMM4236"

    def staff_row(self, **kwargs):
        defaults = dict(
            staff_pf_number=4236, staff_name="Charles Muchiri",
            sales_code=self.CODE, staff_role="SME RM", staff_branch="Rehani",
            active=1, target_loan_disbursement=270_000_000,
            target_new_customers=24, target_deposits_value=120_000_000)
        defaults.update(kwargs)
        return BranchEmployeeDmcData.objects.create(**defaults)

    def branch_row(self, **kwargs):
        defaults = dict(
            staff_pf_number=4236, staff_name="Charles Muchiri",
            sales_code=self.CODE, staff_role="SME RM", staff_branch="Rehani",
            active=1, target_loan_disbursement=14_400_000,
            target_new_customers=4, target_deposits_value=38_400_000,
            target_pbt_revenue=900_000_000)
        defaults.update(kwargs)
        return BranchFinalEmployeeDmcData.objects.create(**defaults)

    def test_the_per_person_row_wins_over_the_branch_row(self):
        self.staff_row()
        self.branch_row()
        targets, table, _ = roster_target_row(self.CODE)
        self.assertEqual(table, STAFF_TABLE)
        self.assertEqual(targets["target_loan_disbursement"], 270_000_000)
        self.assertEqual(targets["target_new_customers"], 24)

    def test_somebody_with_no_personal_row_falls_back_to_the_branch_row(self):
        """A BBM's personal row IS the branch row, so without this they would
        have no targets at all."""
        self.branch_row()
        targets, table, _ = roster_target_row(self.CODE)
        self.assertEqual(table, BRANCH_TABLE)
        self.assertEqual(targets["target_loan_disbursement"], 14_400_000)

    def test_a_branch_only_target_is_not_lent_to_an_rm(self):
        """``target_pbt_revenue`` is on the branch table and not the staff one.

        Handing an RM the BRANCH's revenue target would read as a couple of
        percent and paint a fully performing RM red, so the line reports that
        the column is missing instead of borrowing a figure from a different
        grain."""
        self.staff_row()
        self.branch_row()
        targets, table, columns = roster_target_row(self.CODE)
        self.assertEqual(table, STAFF_TABLE)
        self.assertNotIn("target_pbt_revenue", columns)
        self.assertIsNone(targets.get("target_pbt_revenue"))

        card = build_card(self.CODE)
        lines = [ln for g in card["perspectives"] for ln in g["lines"]]
        profit = [ln for ln in lines if ln["kpi_code"] == "operating_profit"]
        self.assertTrue(profit, "the SME RM card has an Operating Profit line")
        self.assertEqual(profit[0]["pending_label"], "No target column")
        self.assertIsNone(profit[0]["score"])
        self.assertIsNone(profit[0]["ytd_target"])

    def test_the_card_says_which_table_its_targets_came_from(self):
        """A figure nobody can trace is a figure nobody will sign."""
        self.staff_row()
        card = build_card(self.CODE)
        self.assertEqual(card["target_source"], STAFF_TABLE)
        lines = [ln for g in card["perspectives"] for ln in g["lines"]]
        disbursements = [ln for ln in lines
                         if ln["kpi_code"] == "net_asset_disbursements"][0]
        self.assertEqual(disbursements["target_field"],
                         "target_loan_disbursement")
        self.assertEqual(disbursements["annual_target"], 270_000_000)


class ProrationTests(TestCase):
    """The target is sliced to the date the ACTUAL is as at — not to today.

    The eight Q3 cards put the YTD target at 8/12 of the year with actuals to
    31 August. 31 August is day 243 of 365, so the two agree to one part in
    750; comparing an August actual against a target sliced to today would
    hand the RM nine weeks of plan they never had the chance to earn.
    """

    def test_a_month_end_actual_matches_the_cards_own_months_over_twelve(self):
        annual = 270_000_000
        august = datetime.date(2026, 8, 31)
        self.assertAlmostEqual(prorate_to(annual, august),
                               annual * 8 / 12, delta=annual * 0.002)

    def test_the_slice_follows_the_actuals_date_not_today(self):
        annual = 1_200_000
        january = prorate_to(annual, datetime.date(2026, 1, 31))
        december = prorate_to(annual, datetime.date(2026, 12, 31))
        self.assertLess(january, december)
        self.assertAlmostEqual(december, annual, delta=1)

    def test_counts_come_out_whole(self):
        """"589.578 new customers" is not a target anybody can be held to."""
        value = prorate_to(897, datetime.date(2026, 9, 5), counted=True)
        self.assertEqual(value, float(round(value)))

    def test_no_target_stays_no_target(self):
        self.assertIsNone(prorate_to(None, datetime.date(2026, 9, 5)))

    def test_a_monthly_feed_is_dated_to_the_last_closed_month(self):
        """A feed loaded monthly cannot be current to today, and saying it is
        would slice the target past the actual."""
        self.assertEqual(last_closed_month_end(datetime.date(2026, 10, 9)),
                         datetime.date(2026, 9, 30))
        self.assertEqual(last_closed_month_end(datetime.date(2026, 1, 3)),
                         datetime.date(2025, 12, 31))


class LoanLossScoringTests(TestCase):
    """Provisions are a ceiling. Nil provisions is the best outcome, and the
    line on the card is literally titled "Nil provisions"."""

    def test_under_the_ceiling_scores_above_target(self):
        self.assertAlmostEqual(score_for(500_000, 1_000_000, False), 1.2)

    def test_over_the_ceiling_scores_below(self):
        self.assertAlmostEqual(score_for(2_000_000, 1_000_000, False), 0.5)

    def test_nil_provisions_is_the_best_score_not_a_division_by_zero(self):
        self.assertAlmostEqual(score_for(0, 1_000_000, False), 1.2)


class SignoffTests(TestCase):
    """Signing freezes the figures. A signature against numbers that move
    means nothing."""

    CODE = "FJ4145"

    def setUp(self):
        self.user = User.objects.create_user(
            "fjabani", password="x", first_name="Fridah", last_name="Jabani")
        Profile.objects.update_or_create(
            user=self.user,
            defaults={"sales_code": self.CODE, "branch": "Nyeri"})
        BranchEmployeeDmcData.objects.create(
            staff_pf_number=4145, staff_name="Fridah Jabani",
            sales_code=self.CODE, staff_role="SME BBC", staff_branch="Nyeri",
            team_leader="Jane Manager", active=1,
            target_deposits_value=38_400_000, target_new_customers=12,
            target_loan_disbursement=108_000_000)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_signing_stores_the_card_as_it_stood(self):
        response = self.client.post(
            "/staff_management/scorecard-automation/my-card/sign/",
            {"role": "owner", "comment": "Agreed"}, format="json")
        self.assertEqual(response.status_code, 200, response.content)
        signoff = response.data["signoff"]
        self.assertTrue(signoff["signed"])
        self.assertEqual(signoff["signed_by"], "Fridah Jabani")
        self.assertEqual(signoff["owner_comment"], "Agreed")

        from .scorecard_automation.models import ScorecardSignoff
        stored = ScorecardSignoff.objects.get(sales_code=self.CODE)
        self.assertEqual(stored.period, period_for())
        self.assertTrue(stored.card.get("has_card"))
        self.assertEqual(stored.performance_score,
                         stored.card["performance_score"])

    def test_a_signed_card_is_served_as_signed_not_recomputed(self):
        self.client.post(
            "/staff_management/scorecard-automation/my-card/sign/",
            {"role": "owner"}, format="json")
        response = self.client.get(
            "/staff_management/scorecard-automation/my-card/")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.data["live"],
                         "a signed card must not show live figures under the "
                         "signature")
        self.assertTrue(response.data["signoff"]["signed"])

    def test_a_colleague_cannot_counter_sign(self):
        self.client.post(
            "/staff_management/scorecard-automation/my-card/sign/",
            {"role": "owner"}, format="json")
        other = User.objects.create_user("nosy", password="x",
                                         first_name="Not", last_name="Manager")
        client = APIClient()
        client.force_authenticate(other)
        response = client.post(
            "/staff_management/scorecard-automation/my-card/sign/",
            {"role": "manager", "sales_code": self.CODE}, format="json")
        self.assertEqual(response.status_code, 403, response.content)

    def test_the_named_manager_can_counter_sign(self):
        self.client.post(
            "/staff_management/scorecard-automation/my-card/sign/",
            {"role": "owner"}, format="json")
        manager = User.objects.create_user(
            "jmanager", password="x", first_name="Jane", last_name="Manager")
        client = APIClient()
        client.force_authenticate(manager)
        response = client.post(
            "/staff_management/scorecard-automation/my-card/sign/",
            {"role": "manager", "sales_code": self.CODE,
             "comment": "Discussed"}, format="json")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.data["signoff"]["fully_signed"])

    def test_a_manager_cannot_counter_sign_an_unsigned_card(self):
        manager = User.objects.create_user(
            "jmanager2", password="x", first_name="Jane", last_name="Manager")
        client = APIClient()
        client.force_authenticate(manager)
        response = client.post(
            "/staff_management/scorecard-automation/my-card/sign/",
            {"role": "manager", "sales_code": self.CODE}, format="json")
        self.assertEqual(response.status_code, 409, response.content)

    def test_download_returns_a_workbook(self):
        response = self.client.get(
            "/staff_management/scorecard-automation/my-card/download/")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"PK"), "not a .xlsx")

    def test_the_signed_download_carries_the_signature_and_drops_the_warning(self):
        """The file and the signature have to agree: a signed card downloaded
        without the signature on it is a document somebody can present as
        unsigned."""
        from io import BytesIO

        from openpyxl import load_workbook

        self.client.post(
            "/staff_management/scorecard-automation/my-card/sign/",
            {"role": "owner", "comment": "Agreed with my manager"},
            format="json")
        response = self.client.get(
            "/staff_management/scorecard-automation/my-card/download/")
        self.assertEqual(response.status_code, 200, response.content)

        sheet = load_workbook(BytesIO(response.content)).active
        text = "\n".join(
            " ".join("" if c.value is None else str(c.value) for c in row)
            for row in sheet.iter_rows())
        self.assertIn("Fridah Jabani", text)
        self.assertIn("Agreed with my manager", text)
        self.assertNotIn("unsigned", text)

    def test_the_workbook_carries_the_lines_and_a_signature_block(self):
        from openpyxl import load_workbook
        from io import BytesIO

        card = build_card(self.CODE)
        book = load_workbook(BytesIO(build_card_workbook(card).read()))
        sheet = book.active
        text = "\n".join(
            " ".join("" if c.value is None else str(c.value) for c in row)
            for row in sheet.iter_rows())
        self.assertIn("HFC PERFORMANCE SCORECARD", text)
        self.assertIn("Fridah Jabani", text)
        self.assertIn(self.CODE, text)
        self.assertIn("SIGN-OFF", text)
        self.assertIn("Line manager", text)
        self.assertIn("TOTAL PERFORMANCE SCORE", text)
        self.assertIn("unsigned", text,
                      "an unsigned download must say so on the sheet")

    def test_a_pending_line_prints_its_reason_where_the_figure_would_go(self):
        """A pending line has no number, and a zero in the ACTUAL column reads
        as "achieved nothing" - which is a different statement from "nobody
        has this figure". A line that genuinely stands at nil, like drawdowns
        for an RM who has disbursed nothing, is a real zero and stays one."""
        from io import BytesIO

        from openpyxl import load_workbook

        card = build_card(self.CODE)
        lines = {ln["kpi_name"]: ln
                 for g in card["perspectives"] for ln in g["lines"]}
        pending_names = {name for name, ln in lines.items() if ln["pending"]}
        self.assertTrue(pending_names, "nothing is pending, which cannot be "
                                       "right on an empty warehouse")

        book = load_workbook(BytesIO(build_card_workbook(card).read()))
        sheet = book.active
        printed = {}
        for row in sheet.iter_rows():
            if row[0].value and str(row[0].value).isdigit():
                printed[row[1].value] = row[5].value

        for name in pending_names:
            with self.subTest(name):
                self.assertIsInstance(
                    printed.get(name), str,
                    f"{name} is pending but printed {printed.get(name)!r}")


class EveryLineIsAccountedForTests(TestCase):
    """No KPI may reach the card without somebody having decided what it is.

    Without this, adding a role card seeds KPIs that fall through to the
    default and quietly tell the person "measured outside the warehouse" -
    which may be a lie, and is the kind of lie nobody notices because it looks
    like a considered answer.
    """

    def test_every_seeded_kpi_has_an_explicit_source(self):
        from .live_scorecard import SOURCES
        from .scorecard_automation.models import ScKpi

        codes = set(ScKpi.objects.values_list("kpi_code", flat=True))
        self.assertTrue(codes, "the scorecard config was not seeded")
        self.assertEqual(
            sorted(codes - set(SOURCES)), [],
            "these KPIs would fall through to 'measured outside the "
            "warehouse' without anybody deciding that")

    def test_no_source_points_at_a_kpi_that_no_longer_exists(self):
        """A stale entry is a mapping somebody will trust and that never runs."""
        from .live_scorecard import SOURCES
        from .scorecard_automation.models import ScKpi

        codes = set(ScKpi.objects.values_list("kpi_code", flat=True))
        self.assertEqual(sorted(set(SOURCES) - codes), [])

    def test_a_scored_line_names_both_a_target_and_an_actual(self):
        """Half a mapping scores nothing, and reports the wrong reason for it."""
        from .live_scorecard import SOURCES

        for code, source in sorted(SOURCES.items()):
            if source.pending:
                continue
            with self.subTest(code):
                self.assertTrue(source.target_field, f"{code} has no target")
                self.assertTrue(source.actual_key, f"{code} has no actual")
