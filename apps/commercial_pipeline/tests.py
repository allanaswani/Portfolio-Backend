"""Written against what an RM and a Team Leader would notice.

:data:`STORY` is a comment long enough to clear ``MIN_COMMENT_WORDS``. It is
spelled out once and reused, so that a test which is about something else does
not read as though the comment were the point of it.
"""

from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from rest_framework.test import APIClient, APITestCase

from apps.portfolio.models import Profile
from .models import PipelineEntry as P

BASE = "/commercial_book/"

#: Long enough to satisfy the comment minimum from the requirements sheet.
STORY = ("Offer letter signed and returned, valuation instructed, waiting on "
         "the valuer's report before charge preparation.")


def user(username, sales_code="", segment="COMMERCIAL", group=None):
    u = get_user_model().objects.create_user(username=username, password="x")
    Profile.objects.update_or_create(
        user=u, defaults={"sales_code": sales_code, "segment": segment,
                          "branch": "Commercial"})
    if group:
        u.groups.add(Group.objects.get_or_create(name=group)[0])
    return u


class PipelineTests(APITestCase):
    def setUp(self):
        self.rm = user("rm_one", sales_code="CM001")
        self.other_rm = user("rm_two", sales_code="CM002")
        self.tl = user("tl_one", sales_code="CM900", group="tl_portfolio")

        self.mine = P.objects.create(
            kind=P.KIND_ASSET, customer_name="Technomasters Limited",
            sales_code="CM001", rm_name="RM One", segment="COMMERCIAL",
            amount=Decimal("10000000"), amount_to_disburse=Decimal("6000000"),
            product="Contract Finance", broad_stage=P.BROAD_APPROVED)
        self.theirs = P.objects.create(
            kind=P.KIND_ASSET, customer_name="Estina Green Limited",
            sales_code="CM002", rm_name="RM Two", segment="COMMERCIAL",
            amount=Decimal("7500000"))

    def as_(self, who):
        c = APIClient()
        c.force_authenticate(who)
        return c

    # ── who sees what ────────────────────────────────────────────────
    def test_an_rm_sees_only_their_own_lines(self):
        res = self.as_(self.rm).get(f"{BASE}entries/")
        names = [r["customer_name"] for r in res.data["results"]]
        self.assertEqual(names, ["Technomasters Limited"])

    def test_a_team_leader_sees_the_whole_segment(self):
        res = self.as_(self.tl).get(f"{BASE}entries/")
        names = sorted(r["customer_name"] for r in res.data["results"])
        self.assertEqual(names, ["Estina Green Limited", "Technomasters Limited"])

    def test_an_rm_cannot_open_another_rms_line(self):
        res = self.as_(self.rm).get(f"{BASE}entries/{self.theirs.id}/")
        self.assertEqual(res.status_code, 404)

    def test_a_segment_is_matched_through_its_synonyms_not_by_equality(self):
        """COMMERCIAL and LARGE ENTERPRISES are the same segment here."""
        tl = user("tl_large", sales_code="CM901", segment="LARGE ENTERPRISES",
                  group="tl_portfolio")
        res = self.as_(tl).get(f"{BASE}entries/")
        self.assertEqual(res.data["count"], 2)

    # ── the rules the spreadsheet could not enforce ──────────────────
    def test_a_stage_must_belong_to_its_broad_stage(self):
        res = self.as_(self.rm).post(f"{BASE}entries/", {
            "kind": "asset", "customer_name": "Skillman Construction",
            "amount": "30000000", "product": "Contract finance",
            "broad_stage": P.BROAD_DISBURSEMENT, "stage": "at_credit_risk",
            "comments": STORY,
        }, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("stage", res.data)

    def test_a_matching_stage_pair_is_accepted(self):
        res = self.as_(self.rm).post(f"{BASE}entries/", {
            "kind": "asset", "customer_name": "Skillman Construction",
            "amount": "30000000", "product": "Contract finance",
            "broad_stage": P.BROAD_APPLICATION, "stage": "at_credit_risk",
            "comments": STORY,
        }, format="json")
        self.assertEqual(res.status_code, 201, res.data)

    def test_cannot_disburse_more_than_was_requested(self):
        res = self.as_(self.rm).post(f"{BASE}entries/", {
            "kind": "asset", "customer_name": "Almasi Financial",
            "amount": "10000000", "amount_to_disburse": "50000000",
            "product": "Working Capital",
        }, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("amount_to_disburse", res.data)

    def test_a_deposit_needs_a_deposit_type(self):
        res = self.as_(self.rm).post(f"{BASE}entries/", {
            "kind": "deposit", "customer_name": "Modern Precast",
            "amount": "654000000",
        }, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("deposit_product", res.data)

    def test_a_liability_accepts_new_to_bank_as_an_account_number(self):
        """The file has 'New to bank' and 'tba'. Both are real answers."""
        res = self.as_(self.rm).post(f"{BASE}entries/", {
            "kind": "liability", "customer_name": "Sameer Group ltd",
            "amount": "50000000", "account_no": "New to bank",
            "insurance_type": P.INSURANCE_VIC, "comments": STORY,
        }, format="json")
        self.assertEqual(res.status_code, 201, res.data)

    # ── ownership on create ──────────────────────────────────────────
    def test_a_new_line_is_stamped_with_the_rms_own_code(self):
        self.as_(self.rm).post(f"{BASE}entries/", {
            "kind": "trade", "customer_name": "Experian Limited",
            "amount": "50000000", "revenue": "500000", "product": "LC",
            "comments": STORY,
        }, format="json")
        row = P.objects.get(customer_name="Experian Limited")
        self.assertEqual(row.sales_code, "CM001")
        self.assertEqual(row.segment, "COMMERCIAL")

    # ── removing a line ──────────────────────────────────────────────
    def test_deleting_clears_rather_than_erases(self):
        res = self.as_(self.rm).delete(f"{BASE}entries/{self.mine.id}/")
        self.assertEqual(res.status_code, 204)
        self.mine.refresh_from_db()
        self.assertFalse(self.mine.is_active)

        listed = self.as_(self.rm).get(f"{BASE}entries/")
        self.assertEqual(listed.data["count"], 0, "cleared rows are out by default")

        kept = self.as_(self.rm).get(f"{BASE}entries/", {"is_active": "false"})
        self.assertEqual(kept.data["count"], 1, "and can still be looked at")

    # ── what the TL opens it for ─────────────────────────────────────
    def test_the_summary_totals_the_team(self):
        res = self.as_(self.tl).get(f"{BASE}summary/")
        self.assertEqual(res.data["scope"], "team")
        assets = next(r for r in res.data["by_kind"] if r["kind"] == "asset")
        self.assertEqual(assets["count"], 2)
        self.assertEqual(assets["total_amount"], Decimal("17500000"))

    def test_the_summary_for_an_rm_covers_only_their_own(self):
        res = self.as_(self.rm).get(f"{BASE}summary/")
        self.assertEqual(res.data["scope"], "mine")
        assets = next(r for r in res.data["by_kind"] if r["kind"] == "asset")
        self.assertEqual(assets["total_amount"], Decimal("10000000"))

    def test_export_is_scoped_the_same_way_as_the_list(self):
        body = self.as_(self.rm).get(f"{BASE}export/").content.decode()
        self.assertIn("Technomasters Limited", body)
        self.assertNotIn("Estina Green Limited", body)

    def test_options_offers_the_stages_under_each_broad_stage(self):
        res = self.as_(self.rm).get(f"{BASE}options/")
        self.assertIn("at_credit_risk",
                      res.data["stages_under_broad"][P.BROAD_APPLICATION])
        self.assertFalse(res.data["can_see_team"])
        self.assertTrue(self.as_(self.tl).get(f"{BASE}options/").data["can_see_team"])


class WorkbookUploadTests(APITestCase):
    """Loading the workbook without a shell on the host."""

    def setUp(self):
        self.rm = user("up_rm", sales_code="CM010")
        self.tl = user("up_tl", sales_code="CM911", group="tl_portfolio")

    def as_(self, who):
        c = APIClient()
        c.force_authenticate(who)
        return c

    def book(self, name="pipeline.xlsx"):
        """A workbook shaped like the real one, built in memory."""
        import io as _io

        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Assets Pipeline"
        ws.append(["Branch", "Sales Person Name", "Customer Name",
                   "Loan Amount Requested", "Amount to be Disbursed", "Product",
                   "Broad Application stage", "Additional comments"])
        ws.append(["Commercial", "Anne Kimani", "Technomasters Limited",
                   10000000, 6000000, "Contract Finance", "Approved", "Limit in place"])
        # 'Credit Risk' is not a broad stage; the importer must place it properly.
        ws.append(["Commercial", "Irene Njenga", "Cloudarc Solutions Ltd",
                   10000000, 10000000, "Contract Financing", "Credit Risk", ""])
        # A row with no customer is skipped, not imported blank.
        ws.append(["Commercial", "Nobody", "", 500, 500, "x", "Approved", ""])

        trade = wb.create_sheet("Trade Pipeline")
        trade.append(["Branch", "Sales Person Name", "Customer Name", "Amount ",
                      "Revenues", "Type", "Application Stage", "Comments"])
        trade.append(["Commercial", "Lilian Biwott", "Experian Limited",
                      50000000, 500000, "LC", "Approved", "Issued."])

        buf = _io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile(
            name, buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def post(self, who, **extra):
        payload = {"file": self.book()}
        payload.update(extra)
        return self.as_(who).post(f"{BASE}upload/", payload, format="multipart")

    def test_an_rm_cannot_load_everybody_elses_rows(self):
        self.assertEqual(self.post(self.rm).status_code, 403)

    def test_the_first_upload_reports_and_writes_nothing(self):
        res = self.post(self.tl)
        self.assertEqual(res.status_code, 200, res.data)
        self.assertFalse(res.data["applied"])
        self.assertEqual(res.data["total"], 3)
        self.assertEqual(res.data["skipped"], 1, "the row with no customer")
        self.assertEqual(P.objects.count(), 0, "nothing saved on a dry run")

    def test_applying_saves_and_places_the_stages_properly(self):
        res = self.post(self.tl, apply="true")
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(P.objects.count(), 3)

        row = P.objects.get(customer_name="Cloudarc Solutions Ltd")
        self.assertEqual(row.broad_stage, P.BROAD_APPLICATION,
                         "'Credit Risk' is a specific stage, not a broad one")
        self.assertEqual(row.stage, "at_credit_risk")

    def test_replace_removes_imported_rows_but_never_somebody_s_own(self):
        self.post(self.tl, apply="true")
        typed = P.objects.create(kind=P.KIND_ASSET, customer_name="Typed by an RM",
                                 segment="COMMERCIAL", amount=1, created_by=self.rm)

        self.post(self.tl, apply="true", replace="true")
        self.assertTrue(P.objects.filter(pk=typed.pk).exists(),
                        "a row an RM typed must survive a re-upload")
        self.assertEqual(P.objects.filter(created_by__isnull=True).count(), 3,
                         "the previous import was replaced, not doubled")

    def test_a_file_that_is_not_a_workbook_is_refused(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        res = self.as_(self.tl).post(
            f"{BASE}upload/",
            {"file": SimpleUploadedFile("notes.txt", b"hello", content_type="text/plain")},
            format="multipart")
        self.assertEqual(res.status_code, 400)
        self.assertIn("Excel", str(res.data))


class ScopeTests(APITestCase):
    """Who sees whose lines.

    ``portfolio_mgt`` is the group every Commercial RM is in - it is what
    ``lib/roleNavConfig.tsx`` hangs the "My Pipeline" link off, and what
    ``staff_management.targets_views.RM_GROUP`` calls the RM group. It was in
    ``views.TEAM_GROUPS`` when this module shipped, so every RM who could open
    the page was served the whole segment. The tests written at the time did
    not catch it because they built their RMs with no group at all.
    """

    def setUp(self):
        self.rm = user("scope_rm", sales_code="CM101", group="portfolio_mgt")
        self.tl = user("scope_tl", sales_code="CM900", group="tl_portfolio")

        self.mine = P.objects.create(
            kind=P.KIND_ASSET, customer_name="Mine Limited", sales_code="CM101",
            rm_name="Scope RM", segment="COMMERCIAL", amount=1_000_000,
            product="Contract finance")
        self.theirs = P.objects.create(
            kind=P.KIND_ASSET, customer_name="Theirs Limited", sales_code="CM102",
            rm_name="Other RM", segment="COMMERCIAL", amount=2_000_000,
            product="Contract finance")

    def as_(self, who):
        c = APIClient()
        c.force_authenticate(who)
        return c

    def test_an_rm_in_the_rm_group_still_sees_only_their_own_lines(self):
        res = self.as_(self.rm).get(f"{BASE}entries/")
        self.assertEqual([r["customer_name"] for r in res.data["results"]],
                         ["Mine Limited"])

    def test_an_rm_in_the_rm_group_cannot_open_another_rms_line(self):
        res = self.as_(self.rm).get(f"{BASE}entries/{self.theirs.id}/")
        self.assertEqual(res.status_code, 404)

    def test_an_rm_in_the_rm_group_cannot_load_the_workbook(self):
        res = self.as_(self.rm).post(f"{BASE}upload/", {}, format="multipart")
        self.assertEqual(res.status_code, 403)

    def test_being_django_staff_is_not_a_team_role(self):
        """migrate_legacy_auth copies is_staff over verbatim from the old
        system, so it says where an account came from, not what it may read."""
        staffer = user("scope_staff", sales_code="CM103")
        staffer.is_staff = True
        staffer.save(update_fields=["is_staff"])
        res = self.as_(staffer).get(f"{BASE}entries/")
        self.assertEqual(res.data["count"], 0, "no rows carry CM103")

    def test_a_superuser_still_sees_the_team(self):
        boss = user("scope_boss", sales_code="")
        boss.is_superuser = True
        boss.save(update_fields=["is_superuser"])
        self.assertEqual(self.as_(boss).get(f"{BASE}entries/").data["count"], 2)

    # ── scope=mine, which the RM page sends ──
    def test_a_team_leader_asking_for_their_own_book_gets_only_that(self):
        P.objects.create(kind=P.KIND_ASSET, customer_name="The TL's own",
                         sales_code="CM900", segment="COMMERCIAL", amount=5)
        res = self.as_(self.tl).get(f"{BASE}entries/", {"scope": "mine"})
        self.assertEqual([r["customer_name"] for r in res.data["results"]],
                         ["The TL's own"])
        wide = self.as_(self.tl).get(f"{BASE}entries/")
        self.assertEqual(wide.data["count"], 3, "and the team view is untouched")

    def test_the_summary_says_mine_when_mine_was_asked_for(self):
        res = self.as_(self.tl).get(f"{BASE}summary/", {"scope": "mine"})
        self.assertEqual(res.data["scope"], "mine")

    def test_options_does_not_offer_the_owner_picker_on_the_rm_page(self):
        res = self.as_(self.tl).get(f"{BASE}options/", {"scope": "mine"})
        self.assertFalse(res.data["can_see_team"])

    def test_the_export_follows_scope_mine_too(self):
        body = self.as_(self.tl).get(f"{BASE}export/", {"scope": "mine"}).content.decode()
        self.assertNotIn("Mine Limited", body)
        self.assertNotIn("Theirs Limited", body)

    # ── writing into somebody else's book ──
    def test_an_rm_cannot_post_a_line_into_a_colleagues_book(self):
        self.as_(self.rm).post(f"{BASE}entries/", {
            "kind": "asset", "customer_name": "Sneaked In", "amount": "9",
            "product": "Contract finance", "sales_code": "CM102",
            "rm_name": "Other RM", "comments": STORY,
        }, format="json")
        row = P.objects.get(customer_name="Sneaked In")
        self.assertEqual(row.sales_code, "CM101", "stamped with their own code")

    def test_an_rm_cannot_hand_their_line_to_somebody_else(self):
        self.as_(self.rm).patch(f"{BASE}entries/{self.mine.id}/",
                                {"sales_code": "CM102"}, format="json")
        self.mine.refresh_from_db()
        self.assertEqual(self.mine.sales_code, "CM101")

    def test_a_team_leader_may_raise_a_line_on_somebodys_behalf(self):
        self.as_(self.tl).post(f"{BASE}entries/", {
            "kind": "asset", "customer_name": "Raised For An RM", "amount": "9",
            "product": "Contract finance", "sales_code": "CM101",
            "rm_name": "Scope RM", "comments": STORY,
        }, format="json")
        self.assertEqual(P.objects.get(customer_name="Raised For An RM").sales_code,
                         "CM101")


class RequirementsSheetTests(APITestCase):
    """"Additional System requirements.xlsx", from the Commercial RM."""

    def setUp(self):
        self.rm = user("req_rm", sales_code="CM201")

    def as_(self, who):
        c = APIClient()
        c.force_authenticate(who)
        return c

    def post(self, **body):
        return self.as_(self.rm).post(f"{BASE}entries/", body, format="json")

    # ── the names ──
    def test_the_two_pipelines_were_renamed(self):
        labels = dict(P.KIND)
        self.assertEqual(labels[P.KIND_DEPOSIT], "Deposits Pipeline")
        self.assertEqual(labels[P.KIND_LIABILITY], "Insurance Pipeline")

    def test_the_stored_kind_values_did_not_change(self):
        """A rename of a label must not become a silent API break."""
        self.assertEqual(
            sorted(v for v, _ in P.KIND),
            ["asset", "deposit", "liability", "trade"])

    # ── VIC / non-VIC ──
    def test_an_insurance_line_must_say_vic_or_non_vic(self):
        res = self.post(kind="liability", customer_name="Sameer Group ltd",
                        amount="50000000", account_no="New to bank",
                        comments=STORY)
        self.assertEqual(res.status_code, 400)
        self.assertIn("insurance_type", res.data)

    def test_non_vic_is_accepted_and_reads_back_with_its_label(self):
        res = self.post(kind="liability", customer_name="Almasi Financial",
                        amount="1000000", account_no="0123456789",
                        insurance_type=P.INSURANCE_NON_VIC, comments=STORY)
        self.assertEqual(res.status_code, 201, res.data)
        self.assertIn("Non-VIC", res.data["insurance_type_display"])

    def test_an_existing_insurance_line_keeps_working_while_it_is_blank(self):
        """Rows that predate the field. Requiring it retrospectively would
        leave the RM unable to correct an amount."""
        row = P.objects.create(kind=P.KIND_LIABILITY, customer_name="Older line",
                               sales_code="CM201", segment="COMMERCIAL",
                               amount=10, account_no="tba")
        res = self.as_(self.rm).patch(f"{BASE}entries/{row.id}/",
                                      {"amount": "20"}, format="json")
        self.assertEqual(res.status_code, 200, res.data)

    def test_the_insurance_type_cannot_be_blanked_once_set(self):
        row = P.objects.create(kind=P.KIND_LIABILITY, customer_name="Set line",
                               sales_code="CM201", segment="COMMERCIAL",
                               amount=10, account_no="tba",
                               insurance_type=P.INSURANCE_VIC)
        res = self.as_(self.rm).patch(f"{BASE}entries/{row.id}/",
                                      {"insurance_type": ""}, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("insurance_type", res.data)

    # ── the deposit types ──
    def test_cash_margin_and_escrow_are_deposit_types(self):
        for value in ("CASH_MARGIN", "ESCROW"):
            with self.subTest(value):
                res = self.post(kind="deposit", customer_name=f"Depositor {value}",
                                amount="1000000", deposit_product=value,
                                comments=STORY)
                self.assertEqual(res.status_code, 201, res.data)

    def test_cash_margin_fits_in_the_column(self):
        """It is eleven characters; the column held eight when this shipped."""
        self.assertGreaterEqual(
            P._meta.get_field("deposit_product").max_length, len("CASH_MARGIN"))

    # ── the comment minimum ──
    def test_a_new_line_needs_a_comment_that_says_something(self):
        res = self.post(kind="asset", customer_name="Terse Limited",
                        amount="1000000", product="Contract finance",
                        comments="waiting")
        self.assertEqual(res.status_code, 400)
        self.assertIn("comments", res.data)
        self.assertIn(str(P.MIN_COMMENT_WORDS), str(res.data["comments"]))

    def test_moving_the_stage_needs_the_comment_to_keep_up(self):
        row = P.objects.create(kind=P.KIND_ASSET, customer_name="Imported line",
                               sales_code="CM201", segment="COMMERCIAL",
                               amount=10, product="Contract finance",
                               comments="ok")
        res = self.as_(self.rm).patch(
            f"{BASE}entries/{row.id}/",
            {"broad_stage": P.BROAD_CREDIT_EVALUATION, "stage": "ccm"},
            format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("comments", res.data)

        ok = self.as_(self.rm).patch(
            f"{BASE}entries/{row.id}/",
            {"broad_stage": P.BROAD_CREDIT_EVALUATION, "stage": "ccm",
             "comments": STORY}, format="json")
        self.assertEqual(ok.status_code, 200, ok.data)

    def test_an_edit_that_changes_neither_leaves_a_short_comment_alone(self):
        """Otherwise every imported row is frozen until somebody writes it a
        paragraph, and the RM goes back to the spreadsheet."""
        row = P.objects.create(kind=P.KIND_ASSET, customer_name="Imported line",
                               sales_code="CM201", segment="COMMERCIAL",
                               amount=10, product="Contract finance",
                               comments="ok")
        res = self.as_(self.rm).patch(f"{BASE}entries/{row.id}/",
                                      {"amount": "20"}, format="json")
        self.assertEqual(res.status_code, 200, res.data)

    # ── the new stages ──
    def test_the_sheets_stages_are_all_choosable_with_their_unit(self):
        cases = [
            (P.BROAD_RM_ONLY, "relationship_manager"),
            (P.BROAD_CREDIT_ANALYSIS, "credit_origination_manager"),
            (P.BROAD_CREDIT_EVALUATION, "board"),
            (P.BROAD_APPROVED, "instructions_to_lawyers"),
            (P.BROAD_DISBURSEMENT, "trade_middle_officer"),
        ]
        for broad, stage in cases:
            with self.subTest(f"{broad}/{stage}"):
                res = self.post(kind="asset", customer_name=f"Deal {stage}",
                                amount="1000000", product="Contract finance",
                                broad_stage=broad, stage=stage, comments=STORY)
                self.assertEqual(res.status_code, 201, res.data)

    def test_pre_is_valid_at_the_three_points_the_sheet_lists_it(self):
        for broad in (P.BROAD_CREDIT_EVALUATION, P.BROAD_APPROVED,
                      P.BROAD_DISBURSEMENT):
            with self.subTest(broad):
                res = self.post(kind="asset", customer_name=f"PRE at {broad}",
                                amount="1000000", product="Contract finance",
                                broad_stage=broad, stage="pre", comments=STORY)
                self.assertEqual(res.status_code, 201, res.data)

    def test_pre_is_not_valid_where_the_sheet_does_not_list_it(self):
        res = self.post(kind="asset", customer_name="PRE too early",
                        amount="1000000", product="Contract finance",
                        broad_stage=P.BROAD_RM_ONLY, stage="pre", comments=STORY)
        self.assertEqual(res.status_code, 400)
        self.assertIn("stage", res.data)

    # ── Remove-Charge Dispatch to Customer, Remove-Disbursed ──
    def test_a_retired_stage_cannot_be_chosen_for_a_new_line(self):
        res = self.post(kind="asset", customer_name="Old stage",
                        amount="1000000", product="Contract finance",
                        broad_stage=P.BROAD_APPROVED, stage="charge_dispatch",
                        comments=STORY)
        self.assertEqual(res.status_code, 400)
        self.assertIn("stage", res.data)

    def test_the_retired_broad_stage_cannot_be_chosen_either(self):
        res = self.post(kind="asset", customer_name="Old broad stage",
                        amount="1000000", product="Contract finance",
                        broad_stage=P.BROAD_DISBURSED, comments=STORY)
        self.assertEqual(res.status_code, 400)
        self.assertIn("broad_stage", res.data)

    def test_a_line_already_on_a_retired_stage_can_still_be_saved(self):
        """Deleting the choice outright would make these rows unsaveable, and
        an RM who cannot save a row has lost the deal from the pipeline."""
        row = P.objects.create(kind=P.KIND_ASSET, customer_name="Mid-flight",
                               sales_code="CM201", segment="COMMERCIAL",
                               amount=10, product="Contract finance",
                               broad_stage=P.BROAD_APPROVED,
                               stage="charge_dispatch", comments="ok")
        res = self.as_(self.rm).patch(f"{BASE}entries/{row.id}/",
                                      {"amount": "20"}, format="json")
        self.assertEqual(res.status_code, 200, res.data)
        row.refresh_from_db()
        self.assertEqual(row.stage, "charge_dispatch")

    def test_retired_choices_are_not_offered_but_still_have_labels(self):
        res = self.as_(self.rm).get(f"{BASE}options/")
        offered = {s["value"] for s in res.data["stages"]}
        self.assertNotIn("charge_dispatch", offered)
        self.assertIn("charge_dispatch", res.data["stage_labels"],
                      "a stored row still has to render as words")
        broad = {s["value"] for s in res.data["broad_stages"]}
        self.assertNotIn(P.BROAD_DISBURSED, broad)
        self.assertNotIn("charge_dispatch",
                         res.data["stages_under_broad"][P.BROAD_APPROVED])

    def test_options_carries_the_new_choices_and_the_word_minimum(self):
        res = self.as_(self.rm).get(f"{BASE}options/")
        self.assertEqual([i["value"] for i in res.data["insurance_types"]],
                         ["VIC", "NON_VIC"])
        self.assertEqual(res.data["min_comment_words"], P.MIN_COMMENT_WORDS)
        self.assertIn("insurance_type",
                      res.data["fields_by_kind"][P.KIND_LIABILITY])
        self.assertIn("ESCROW", [d["value"] for d in res.data["deposit_products"]])


class VocabularyTests(APITestCase):
    """The stage lists have to agree with each other.

    Three structures describe the same vocabulary - ``STAGE``, ``BROAD_STAGE``
    and ``STAGES_UNDER_BROAD`` - plus the importer's ``STAGE_MAP``. A slug that
    appears in one and not another is a stage that either cannot be chosen or
    cannot be validated, and neither failure is visible until somebody tries.
    """

    def test_every_stage_sits_under_a_broad_stage_that_exists(self):
        broads = dict(P.BROAD_STAGE)
        stages = dict(P.STAGE)
        for broad, members in P.STAGES_UNDER_BROAD.items():
            self.assertIn(broad, broads)
            for stage in members:
                self.assertIn(stage, stages, f"{broad} points at {stage}")

    def test_every_broad_stage_has_its_stages_listed(self):
        for broad, _ in P.BROAD_STAGE:
            self.assertIn(broad, P.STAGES_UNDER_BROAD)

    def test_no_stage_is_left_unreachable(self):
        placed = set().union(*P.STAGES_UNDER_BROAD.values())
        for stage, _ in P.STAGE:
            self.assertIn(stage, placed, f"{stage} belongs to no broad stage")

    def test_a_retired_value_is_still_a_value(self):
        self.assertTrue(P.RETIRED_STAGES <= set(dict(P.STAGE)))
        self.assertTrue(P.RETIRED_BROAD_STAGES <= set(dict(P.BROAD_STAGE)))

    def test_every_choice_fits_its_column(self):
        for field, values in (
            ("stage", [v for v, _ in P.STAGE]),
            ("broad_stage", [v for v, _ in P.BROAD_STAGE]),
            ("deposit_product", [v for v, _ in P.DEPOSIT_PRODUCT]),
            ("insurance_type", [v for v, _ in P.INSURANCE_TYPE]),
        ):
            longest = max(values, key=len)
            self.assertLessEqual(len(longest), P._meta.get_field(field).max_length,
                                 f"{field}: {longest!r} does not fit")

    def test_the_importer_only_maps_onto_values_that_exist_and_validate(self):
        from .importer import DEPOSIT_PRODUCTS, STAGE_MAP

        broads, stages = dict(P.BROAD_STAGE), dict(P.STAGE)
        for raw, (broad, specific) in STAGE_MAP.items():
            if broad is not None:
                self.assertIn(broad, broads, raw)
                self.assertNotIn(broad, P.RETIRED_BROAD_STAGES,
                                 f"{raw} imports onto a retired broad stage")
            if specific:
                self.assertIn(specific, stages, raw)
                self.assertNotIn(specific, P.RETIRED_STAGES,
                                 f"{raw} imports onto a retired stage")
            if broad is not None and specific:
                self.assertIn(specific, P.STAGES_UNDER_BROAD[broad],
                              f"{raw} imports a pair the serializer rejects")
        for raw, value in DEPOSIT_PRODUCTS.items():
            self.assertIn(value, dict(P.DEPOSIT_PRODUCT), raw)

    def test_every_field_a_kind_uses_is_a_real_field(self):
        for kind, fields in P.FIELDS_BY_KIND.items():
            for name in fields:
                P._meta.get_field(name)   # raises if it is not there


class ImporterAfterTheRenameTests(APITestCase):
    """What happens to the next copy of the workbook.

    The two sheets were renamed in the app, so the desk's next save is likely
    to carry the new tab names - and may well carry a VIC column, since that
    is now a thing the pipeline tracks. Both have to work, and a renamed tab
    must not be read as an absent one: that would drop a whole pipeline out of
    the Team Leader's totals with nothing to show it had ever been there.
    """

    def build(self, deposits_tab="Deposits", insurance_tab="Liabilities",
              vic_header=None, vic_value=None, stage="Credit Evaluation",
              deposit_type="CASA"):
        import io as _io

        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Assets Pipeline"
        ws.append(["Branch", "Sales Person Name", "Customer Name",
                   "Loan Amount Requested", "Amount to be Disbursed", "Product",
                   "Broad Application stage", "Additional comments"])
        ws.append(["Commercial", "Anne Kimani", "Technomasters Limited",
                   10000000, 6000000, "Contract Finance", stage, "Limit in place"])

        dep = wb.create_sheet(deposits_tab)
        dep.append(["Branch", "Sales Person Name", "Customer Name",
                    "Receipt date/By When", "Product", "Amount", "Comments"])
        dep.append(["Commercial", "Anne Kimani", "Modern Precast",
                    "30/1/2026", deposit_type, 654000000, "Confirmed"])

        ins = wb.create_sheet(insurance_tab)
        header = ["RM", "Branch", "Customer Name", "Account No", "Amount",
                  "Date expected", "Comments"]
        row = ["Anne Kimani", "Commercial", "Sameer Group ltd", "New to bank",
               50000000, "30/1/2026", "Quote issued"]
        if vic_header:
            header.append(vic_header)
            row.append(vic_value)
        ins.append(header)
        ins.append(row)

        buf = _io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile(
            "pipeline.xlsx", buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def read(self, **kwargs):
        from .importer import parse_workbook

        return parse_workbook(self.build(**kwargs))

    def test_the_original_tab_names_still_work(self):
        entries, warnings, _ = self.read()
        self.assertEqual(len(entries), 3, warnings)

    def test_the_renamed_tabs_are_found_and_said_so(self):
        entries, warnings, _ = self.read(deposits_tab="Deposits Pipeline",
                                         insurance_tab="Insurance Pipeline")
        kinds = sorted(e.kind for e in entries)
        self.assertEqual(kinds, ["asset", "deposit", "liability"])
        self.assertTrue(any("Insurance Pipeline" in w for w in warnings),
                        f"the reader should say which tab it read: {warnings}")

    def test_a_vic_column_is_read_from_wherever_it_is(self):
        entries, warnings, _ = self.read(insurance_tab="Insurance Pipeline",
                                         vic_header="VIC / Non-VIC",
                                         vic_value="Non-VIC")
        row = next(e for e in entries if e.kind == P.KIND_LIABILITY)
        self.assertEqual(row.insurance_type, P.INSURANCE_NON_VIC, warnings)

    def test_vic_is_not_read_off_a_column_merely_containing_those_letters(self):
        """"Service" contains v-i-c. Matching it would file every row under
        whatever that column happens to say."""
        entries, _, _ = self.read(vic_header="Service", vic_value="Branch desk")
        row = next(e for e in entries if e.kind == P.KIND_LIABILITY)
        self.assertEqual(row.insurance_type, "")

    def test_an_unrecognised_vic_value_is_reported_not_guessed(self):
        entries, warnings, _ = self.read(vic_header="VIC", vic_value="maybe")
        row = next(e for e in entries if e.kind == P.KIND_LIABILITY)
        self.assertEqual(row.insurance_type, "")
        self.assertTrue(any("neither VIC nor" in w for w in warnings), warnings)

    def test_the_sheets_new_stage_names_import(self):
        entries, warnings, _ = self.read(stage="Credit Evaluation")
        row = next(e for e in entries if e.kind == P.KIND_ASSET)
        self.assertEqual(row.broad_stage, P.BROAD_CREDIT_EVALUATION, warnings)

    def test_disbursed_lands_on_disbursement_not_on_the_retired_value(self):
        entries, _, _ = self.read(stage="Disbursed")
        row = next(e for e in entries if e.kind == P.KIND_ASSET)
        self.assertEqual(row.broad_stage, P.BROAD_DISBURSEMENT)
        self.assertNotIn(row.stage, P.RETIRED_STAGES)

    def test_a_unit_that_sits_at_several_stages_is_reported_not_placed(self):
        """PRE is listed three times in the requirements sheet. Guessing which
        one a cell means puts the row in the wrong bucket of the totals."""
        entries, warnings, _ = self.read(stage="PRE")
        row = next(e for e in entries if e.kind == P.KIND_ASSET)
        self.assertEqual(row.broad_stage, "")
        self.assertEqual(row.stage, "")
        self.assertTrue(any("more than one stage" in w for w in warnings), warnings)

    def test_cash_margin_imports_as_a_deposit_type(self):
        entries, warnings, _ = self.read(deposit_type="Cash Margin")
        row = next(e for e in entries if e.kind == P.KIND_DEPOSIT)
        self.assertEqual(row.deposit_product, "CASH_MARGIN", warnings)

    def test_an_unknown_deposit_type_is_reported_not_dropped(self):
        entries, warnings, _ = self.read(deposit_type="Call account")
        row = next(e for e in entries if e.kind == P.KIND_DEPOSIT)
        self.assertEqual(row.deposit_product, "")
        self.assertEqual(row.customer_name, "Modern Precast",
                         "the row itself is kept")
        self.assertTrue(any("deposit type" in w for w in warnings), warnings)
